"""2.1.1: `pennywarden self-check-documents` - build a slide deck with speaker notes and a Word file with a header,
footer, styles, settings and comments in memory, the way the reports do, and report whether the bundled templates of
python-pptx and python-docx were found.

Why: those libraries open their XML templates as `<package>/<subfolder>/../templates/<name>.xml`. In the macOS app
the Python code is inside an archive, so `<subfolder>` is not a real folder and macOS can't resolve the `..`
(Windows simplifies the path by its text, which hid this). The fundraiser deck and the Word report failed on a Mac
with FileNotFoundError ...pptx/oxml/../templates/notesMaster.xml. packaging/pyinstaller/pennywarden-mac.spec now
bundles a file in each such subfolder; the macOS smoke test runs this command inside the built app to prove it.
Nothing is written to disk and no data folder is used."""
from __future__ import annotations

import io
import sys
import traceback


def _deck() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "speaker notes"   # notesMaster.xml / notes.xml templates
    prs.save(io.BytesIO())


def _word() -> None:
    from docx import Document

    doc = Document()
    sec = doc.sections[0]
    sec.footer.paragraphs[0].text = "footer"                      # default-footer.xml
    sec.header.paragraphs[0].text = "header"                      # default-header.xml
    doc.settings                                                  # default-settings.xml (when missing)  # noqa: B018
    doc.styles                                                    # default-styles.xml (when missing)    # noqa: B018
    run = doc.add_paragraph("text").add_run("comment")
    doc.add_comment(run, text="comment", author="self-check")    # default-comments.xml
    doc.save(io.BytesIO())


def main(argv: list[str] | None = None) -> int:
    failed = False
    for name, fn in (("slide deck with speaker notes", _deck), ("Word file with header, footer and comments", _word)):
        try:
            fn()
            print(f"OK: {name}")
        except Exception:   # noqa: BLE001 - report every failure, then exit non-zero
            failed = True
            print(f"FAILED: {name}", file=sys.stderr)
            traceback.print_exc()
    return 1 if failed else 0
