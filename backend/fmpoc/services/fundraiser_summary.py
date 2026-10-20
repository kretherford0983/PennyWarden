"""2.1.0 (#275): the fundraiser stakeholder report - a "Financial & Proceeds Statement" for the committee and
stakeholders, as a PDF and as an editable Word document (.docx). It is not audit evidence: no attachments, documents,
transaction or allocation IDs, bank accounts, budgets, excluded lines or filter text. The audit detail stays in
Reports -> Fundraiser detail (audit) and in the Audit / Close reports.

`summary()` computes everything once from the fundraiser's figures (the same rows as the fundraiser page,
`fundraisers._lines()`: excluded lines removed, classified cash float not counted), so the report, the presentation
(#276) and the Fiscal Year report (#277) can't disagree. Bucket nets (charge - less, #278) are the money; charges and
credits are detail rows only.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import tempfile
from fractions import Fraction
from xml.sax.saxutils import escape

import docx  # noqa: F401  - imported at start-up, so a package built without python-docx fails its smoke test
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from ..models import Fundraiser, Workspace, utcnow
from ..money import fmt
from . import fundraisers as fsvc
from .reports import _FONT, _FONT_BOLD, FRAME_W, PM, P, S, _AuditDoc, _Mark, _stamp_and_write, money

TWO_COLUMN_AFTER = 12
STATUS_TEXT = {"INTERIM": "Interim, figures may change", "POST_EVENT": "Post-event", "FINAL": "Final",
               "CANCELLED": "Cancelled"}


# ------------------------------------------------------------------ figures
def pct(part: int, whole: int) -> str | None:
    if whole == 0:
        return None
    v = Fraction(part * 1000, whole)
    n = int(abs(v) + Fraction(1, 2))
    n = n if v >= 0 else -n
    return f"{'-' if n < 0 else ''}{abs(n) // 10}.{abs(n) % 10}"


def report_status(db: Session, f: Fundraiser) -> dict:
    """Interim (planned / in progress), Post-event (ended), Final (marked by a Budget Manager, or every Fiscal Year
    the fundraiser uses is Closed - whichever comes first), Cancelled (with the reason)."""
    fys = fsvc._fys_of(db, f)
    all_closed = bool(fys) and all(y.status == "CLOSED" for y in fys)
    final = f.results_final_at is not None or all_closed
    if f.cancelled_at is not None:
        code = "CANCELLED"
    elif final:
        code = "FINAL"
    else:
        code = "INTERIM" if fsvc.status_of(f) in ("PLANNED", "IN_PROGRESS") else "POST_EVENT"
    return {"code": code, "text": STATUS_TEXT[code], "final": final, "marked_final": f.results_final_at is not None,
            "final_because_closed": all_closed and f.results_final_at is None,
            "marked_final_at": f.results_final_at.isoformat() + "Z" if f.results_final_at else None,
            "cancel_reason": f.cancel_reason if f.cancelled_at else None}


def summary(db: Session, ctx, f: Fundraiser, *, itemized: bool = True, names: bool = True,
            plain_lines: bool = False) -> dict:
    ws = db.get(Workspace, ctx.workspace_id)
    rows, _skipped = fsvc._lines(db, f)
    rows = [r for r in rows if not r["excluded"]]
    inc_rows = [r for r in rows if r["kind"] == "INCOME"]
    exp_rows = [r for r in rows if r["kind"] == "EXPENSE"]
    revenue = sum(r["eff"] for r in inc_rows)
    expenses = sum(r["eff"] for r in exp_rows)
    net = revenue - expenses
    gross = sum(r["a"].amount_cents for r in inc_rows)
    float_returned = sum(r["classified"] for r in inc_rows)
    float_out = sum(r["classified"] for r in exp_rows)

    # contributor / vendor labels: the Entity (not system ones), else the line description. Without names every
    # Entity becomes "Contributor n" (income) or the line description / "Expense" (expenses).
    numbers: dict[tuple, int] = {}

    def who(r) -> str:
        a, t = r["a"], r["t"]
        ent = a.entity or t.parent_entity
        real = ent is not None and not ent.is_system
        if r["kind"] == "INCOME":
            if names:
                return ent.display_name if real else (a.description or "—")
            key = ("e", ent.id) if real else ("l", a.id)
            numbers.setdefault(key, len(numbers) + 1)
            return f"Contributor {numbers[key]}"
        if names and real:
            return " — ".join(v for v in (ent.display_name, a.description) if v)
        return a.description or "Expense"

    def line(r, amount: int, bl=None) -> dict:
        w = who(r)
        note = r["a"].description if r["kind"] == "INCOME" and r["a"].description != w else None
        d = {"date": r["t"].transaction_date.isoformat(), "who": w, "amount": fmt(amount), "note": note, "charge": None, "less": None,
             "less_description": None}
        if bl is not None and (bl.less_cents or 0) > 0:
            d.update(charge=fmt(bl.charge_cents if bl.charge_cents is not None else bl.amount_cents),
                     less=fmt(bl.less_cents), less_description=bl.less_description)
        return d

    buckets = fsvc.ordered_buckets(db, f)
    cats = fsvc.ordered_categories(db, f)
    has_buckets = bool(buckets)
    per_bucket = {b.id: {"INCOME": [], "EXPENSE": []} for b in buckets}
    other = {"INCOME": [], "EXPENSE": []}
    for r in rows:
        for bl in r["buckets"]:
            if bl.bucket_id in per_bucket:
                per_bucket[bl.bucket_id][r["kind"]].append(line(r, bl.amount_cents, bl))
        if r["unassigned"] > 0:
            other[r["kind"]].append(line(r, r["unassigned"]))

    def cents(lines: list[dict]) -> int:
        return sum(int(round(float(x["amount"]) * 100)) for x in lines)

    def bucket_out(b) -> dict:
        i, e = per_bucket[b.id]["INCOME"], per_bucket[b.id]["EXPENSE"]
        ic, ec = cents(i), cents(e)
        return {"id": b.id, "name": b.name, "income_lines": i, "expense_lines": e, "income": fmt(ic),
                "expense": fmt(ec), "net": fmt(ic - ec), "income_pct": pct(ic, revenue), "expense_pct": pct(ec, expenses),
                "credits": fmt(sum(int(round(float(x["less"]) * 100)) for x in e if x["less"]))}

    groups = []
    cat_ids = {c.id for c in cats}
    for c in cats:
        bs = [bucket_out(b) for b in buckets if b.category_id == c.id]
        groups.append({"kind": "category", "name": c.name, "buckets": bs})
    for b in buckets:
        if b.category_id not in cat_ids:
            groups.append({"kind": "bucket", "name": b.name, "buckets": [bucket_out(b)]})
    for g in groups:
        ic = sum(int(round(float(b["income"]) * 100)) for b in g["buckets"])
        ec = sum(int(round(float(b["expense"]) * 100)) for b in g["buckets"])
        g.update(income=fmt(ic), expense=fmt(ec), net=fmt(ic - ec), income_pct=pct(ic, revenue),
                 expense_pct=pct(ec, expenses), income_lines=sum(len(b["income_lines"]) for b in g["buckets"]),
                 expense_count=sum(len(b["expense_lines"]) for b in g["buckets"]))
    oi, oe = cents(other["INCOME"]), cents(other["EXPENSE"])
    other_out = {"name": "Other (not in a bucket)", "income_lines": other["INCOME"], "expense_lines": other["EXPENSE"],
                 "income": fmt(oi), "expense": fmt(oe), "net": fmt(oi - oe), "income_pct": pct(oi, revenue),
                 "expense_pct": pct(oe, expenses)}
    credits = sum(int(round(float(b["credits"]) * 100)) for g in groups for b in g["buckets"])

    fys = fsvc._fys_of(db, f)
    per_fy = []
    if len(fys) > 1:
        for y in fys:
            sub = [r for r in rows if r["fiscal_year_id"] == y.id]
            i = sum(r["eff"] for r in sub if r["kind"] == "INCOME")
            e = sum(r["eff"] for r in sub if r["kind"] == "EXPENSE")
            per_fy.append({"fiscal_year": y.display_name, "income": fmt(i), "expense": fmt(e), "net": fmt(i - e)})

    # key takeaways: factual statements, each only when it applies
    take = []
    margin = pct(net, revenue)
    if revenue:
        take.append(f"Net proceeds of {money(fmt(net))}, a net margin of {margin}% of net revenue.")
    elif net or expenses:
        take.append(f"Net proceeds of {money(fmt(net))}.")
    if has_buckets and revenue:
        top = max(groups, key=lambda g: float(g["income"]), default=None)
        if top is not None and float(top["income"]) > 0:
            take.append(f"{top['name']} was the largest source of revenue: {money(top['income'])} "
                        f"({top['income_pct']}% of net revenue).")
        covers = [g for g in groups if expenses and float(g["income"]) * 100 >= expenses]
        if covers:
            names_ = ", ".join(g["name"] for g in covers)
            take.append(f"{names_} income alone covered all expenses ({money(fmt(expenses))})." if len(covers) == 1
                        else f"Each of {names_} covered all expenses ({money(fmt(expenses))}) on its own.")
    if float_out or float_returned:
        var = float_returned - float_out
        take.append(f"Cash float: {money(fmt(float_out))} taken out, {money(fmt(float_returned))} returned "
                    + ("(fully returned)." if var == 0 else f"(variance {money(fmt(var))})."))
    if credits:
        take.append(f"Credits received on expenses: {money(fmt(credits))}, taken off the charges shown in the "
                    "expense breakdown.")

    plain = []
    if plain_lines and not has_buckets:
        for r in rows:
            plain.append({"date": r["t"].transaction_date.isoformat(), "who": who(r),
                          "income": fmt(r["eff"]) if r["kind"] == "INCOME" else None,
                          "expense": fmt(r["eff"]) if r["kind"] == "EXPENSE" else None})
    status = report_status(db, f)
    return {
        "organization": ws.name, "fundraiser": f.name, "start_date": f.start_date.isoformat(),
        "end_date": f.end_date.isoformat(), "status": status,
        "options": {"itemized": itemized, "names": names, "plain_lines": plain_lines},
        "totals": {"gross_deposits": fmt(gross), "cash_float_returned": fmt(float_returned),
                   "cash_float_out": fmt(float_out), "net_revenue": fmt(revenue), "expenses": fmt(expenses),
                   "net_proceeds": fmt(net), "net_margin": margin, "roi": fsvc.roi(revenue, expenses),
                   "credits": fmt(credits)},
        "per_fiscal_year": per_fy, "has_buckets": has_buckets, "groups": groups, "other": other_out,
        "takeaways": take, "summary_notes": f.summary_notes, "plain_lines": plain,
        "generated_at": utcnow().isoformat(timespec="seconds") + "Z",
    }


def _us(iso: str) -> str:
    d = dt.date.fromisoformat(iso[:10])
    return f"{d.month}/{d.day}/{d.year}"


def event_dates(s: dict) -> str:
    return _us(s["start_date"]) if s["start_date"] == s["end_date"] else f"{_us(s['start_date'])} – {_us(s['end_date'])}"


def _p(v: str | None) -> str:
    return "n/a" if v is None else f"{v}%"


def safe_name(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "-" for c in name).strip("-")[:60] or "fundraiser"


# ------------------------------------------------------------------ PDF
_BAND = colors.HexColor("#1f3a5f")
_TILE = colors.HexColor("#eef3f9")
_W = ParagraphStyle("fs_w", parent=S["body"], textColor=colors.white)
_WT = ParagraphStyle("fs_wt", parent=S["h1"], textColor=colors.white, fontSize=18, leading=22, spaceAfter=0)
_WB = ParagraphStyle("fs_wb", parent=_W, fontName=_FONT_BOLD)
_TILE_V = ParagraphStyle("fs_tv", parent=S["body"], fontName=_FONT_BOLD, fontSize=17, leading=21)
_TILE_K = ParagraphStyle("fs_tk", parent=S["body"], fontName=_FONT_BOLD, fontSize=9, leading=11)
_IND = ParagraphStyle("fs_ind", parent=S["cell"], leftIndent=10)
_IND2 = ParagraphStyle("fs_ind2", parent=S["cell"], leftIndent=20)
_RB = ParagraphStyle("fs_rb", parent=S["cellr"], fontName=_FONT_BOLD)
S.update({"fs_wb": _WB, "fs_rb": _RB, "fs_w": _W, "fs_tv": _TILE_V, "fs_tk": _TILE_K, "fs_ind": _IND, "fs_ind2": _IND2})


def _table(data, widths, header=True, bold_rows=(), italic_rows=()):
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    st = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c3cad3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
          ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    if header:
        st.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe7f1")))
    for r in bold_rows:
        st.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f2f5f9")))
    t.setStyle(TableStyle(st))
    return t


def _h(cells):
    return [P(c, "fs_rb" if i and c in ("Amount", "Subtotal", "% of revenue", "% of expenses",
                                                                  "Income", "Expenses", "Net", "Lines") else "cellb")
            for i, c in enumerate(cells)]


def _b(text, style="cellb"):
    return P(text, style)


def expense_line_rows(x: dict) -> list[tuple[str, str]]:
    """An expense line in the breakdown: with a credit (#278) the charge, "Less: <description>" and the net."""
    if not x.get("less"):
        return [(x["who"], money(x["amount"]))]
    return [(f"{x['who']} — charge", money(x["charge"])), (f"Less: {x['less_description'] or ''}", f"-{money(x['less'])}"),
            ("Net", money(x["amount"]))]


def _charge_rows(x: dict, ncols: int, amount_col: int) -> list[list]:
    """Charge / Less detail under a line: the charge, "Less: <description>" and the net (the line's amount)."""
    if not x.get("less"):
        return []
    blank = [""] * ncols
    a = list(blank)
    a[0] = P("Charge", "fs_ind2")
    a[amount_col] = P(money(x["charge"]), "cellr")
    b = list(blank)
    b[0] = P(f"Less: {x['less_description'] or ''}", "fs_ind2")
    b[amount_col] = P(f"-{money(x['less'])}", "cellr")
    return [a, b]


def build_pdf(s: dict) -> tuple[bytes, _AuditDoc]:
    buf = io.BytesIO()
    doc = _AuditDoc(buf, title="Financial & Proceeds Statement", author=s["organization"])
    W = FRAME_W - 12
    t = s["totals"]
    st = s["status"]
    band = Table([[[P(s["organization"], "fs_w"), Paragraph(escape(s["fundraiser"]), _WT),
                    P("Financial & Proceeds Statement", "fs_wb"),
                    P(f"{event_dates(s)}  ·  {st['text']}", "fs_w")]]], colWidths=[W])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _BAND), ("LEFTPADDING", (0, 0), (-1, -1), 12),
                              ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    f: list = [_Mark(doc, "Transaction-level evidence is in the Audit report"), band, Spacer(1, 8)]
    if st["code"] == "CANCELLED":
        warn = Table([[P(f"Cancelled: {st['cancel_reason'] or 'no reason given'}. The figures below are what was "
                         "received and spent.", "body")]], colWidths=[W])
        warn.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#b3261e")),
                                  ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fde4e2"))]))
        f += [warn, Spacer(1, 8)]
    # executive summary tiles
    tiles = [("Total net revenue", t["net_revenue"], "All income counted for the fundraiser"
              + ("; cash float returned excluded" if t["cash_float_returned"] != "0.00" else "")),
             ("Total expenses", t["expenses"], "Fundraiser expenses" + ("; cash float not included"
                                                                       if t["cash_float_out"] != "0.00" else "")),
             ("Net proceeds", t["net_proceeds"], f"Net margin {_p(t['net_margin'])} of net revenue")]
    cells = [[[P(k, "fs_tk"), PM(escape(money(v)), "fs_tv"), P(c, "small")] for k, v, c in tiles]]
    tt = Table(cells, colWidths=[W / 3] * 3)
    tt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _TILE), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#c3cad3")),
                            ("LINEAFTER", (0, 0), (1, 0), 0.5, colors.HexColor("#c3cad3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    f += [PM("Executive summary", "h2"), tt, Spacer(1, 8)]
    # financial summary
    data = [_h(["Component", "Details", "Amount"])]
    data.append([P("Gross deposits", "cell"), P("All deposits counted for the fundraiser", "cell"), P(money(t["gross_deposits"]), "cellr")])
    if t["cash_float_returned"] != "0.00":
        data.append([P("Less cash float returned", "cell"), P("Cash float coming back inside deposits; not income", "cell"),
                     P(f"-{money(t['cash_float_returned'])}", "cellr")])
    data.append([_b("Net revenue"), P("", "cell"), P(money(t['net_revenue']), "fs_rb")])
    data.append([P("Less expenses", "cell"), P("Fundraiser expenses" + ("; cash float taken out is not an expense"
                                                                         if t["cash_float_out"] != "0.00" else ""), "cell"),
                 P(f"-{money(t['expenses'])}", "cellr")])
    data.append([_b("Net proceeds"), P(f"Net margin {_p(t['net_margin'])}", "cell"), P(money(t['net_proceeds']), "fs_rb")])
    f += [PM("Financial summary", "h2"), _table(data, [1.8 * inch, W - 3.0 * inch, 1.2 * inch], bold_rows=[len(data) - 1])]
    f.append(P(f"ROI (net proceeds ÷ expenses): {t['roi'] if t['roi'] is not None else 'n/a (no expenses)'}", "small"))
    for y in s["per_fiscal_year"]:
        f.append(P(f"{y['fiscal_year']}: revenue {money(y['income'])}, expenses {money(y['expense'])}, "
                   f"net {money(y['net'])}", "small"))
    groups, other = s["groups"], s["other"]
    if s["has_buckets"]:
        # revenue breakdown
        data, bold = [_h(["Category / bucket", "Lines", "Subtotal", "% of revenue"])], []
        for g in groups:
            if g["income"] == "0.00":
                continue
            data.append([_b(g["name"]), P(str(g["income_lines"]), "cellr"), P(money(g['income']), "fs_rb"),
                         P(_p(g["income_pct"]), "cellr")])
            bold.append(len(data) - 1)
            if g["kind"] == "category":
                for b in g["buckets"]:
                    if b["income"] != "0.00":
                        data.append([P(b["name"], "fs_ind"), P(str(len(b["income_lines"])), "cellr"),
                                     P(money(b["income"]), "cellr"), P(_p(b["income_pct"]), "cellr")])
        if other["income"] != "0.00":
            data.append([P(other["name"], "cell"), P(str(len(other["income_lines"])), "cellr"),
                         P(money(other["income"]), "cellr"), P(_p(other["income_pct"]), "cellr")])
        data.append([_b("Total net revenue"), "", P(money(t['net_revenue']), "fs_rb"), P("100.0%" if t["net_revenue"] != "0.00" else "n/a", "cellr")])
        f += [Spacer(1, 8), PM("Revenue breakdown", "h2"), _table(data, [W - 3.1 * inch, 0.7 * inch, 1.2 * inch, 1.2 * inch], bold_rows=bold)]
        # itemized revenue
        if s["options"]["itemized"]:
            f += [Spacer(1, 8), PM("Itemized revenue", "h2")]
            for g in groups:
                for b in g["buckets"]:
                    if b["income_lines"]:
                        title = b["name"] if g["kind"] == "bucket" else f"{g['name']} — {b['name']}"
                        f += _itemized(title, b["income_lines"], b["income"], W)
            if other["income_lines"]:
                f += _itemized(other["name"], other["income_lines"], other["income"], W)
        # expense breakdown
        data, bold = [_h(["Category / bucket / line", "Amount", "% of expenses"])], []
        for g in groups:
            if g["expense"] == "0.00":
                continue
            data.append([_b(g["name"]), P(money(g['expense']), "fs_rb"), P(_p(g["expense_pct"]), "cellr")])
            bold.append(len(data) - 1)
            for b in g["buckets"]:
                if not b["expense_lines"]:
                    continue
                if g["kind"] == "category":
                    data.append([P(b["name"], "fs_ind"), P(money(b["expense"]), "cellr"), P(_p(b["expense_pct"]), "cellr")])
                for x in b["expense_lines"]:
                    for label, amt in expense_line_rows(x):
                        data.append([P(label, "fs_ind2"), P(amt, "cellr"), ""])
        if other["expense_lines"]:
            data.append([_b(other["name"]), P(money(other['expense']), "fs_rb"), P(_p(other["expense_pct"]), "cellr")])
            bold.append(len(data) - 1)
            for x in other["expense_lines"]:
                data.append([P(x["who"], "fs_ind2"), P(money(x["amount"]), "cellr"), ""])
        data.append([_b("Total expenses"), P(money(t['expenses']), "fs_rb"), P("100.0%" if t["expenses"] != "0.00" else "n/a", "cellr")])
        f += [Spacer(1, 8), PM("Expense breakdown", "h2"), _table(data, [W - 2.4 * inch, 1.2 * inch, 1.2 * inch], bold_rows=bold)]
        if t["cash_float_out"] != "0.00":
            f.append(P(f"Cash float taken out ({money(t['cash_float_out'])}) is not an expense; "
                       f"{money(t['cash_float_returned'])} came back in the deposits.", "small"))
        # net by category / bucket
        data, bold = [_h(["Category / bucket", "Income", "Expenses", "Net"])], []
        for g in groups:
            data.append([_b(g["name"]), P(money(g['income']), "fs_rb"), P(money(g['expense']), "fs_rb"),
                         P(money(g['net']), "fs_rb")])
            bold.append(len(data) - 1)
            if g["kind"] == "category":
                for b in g["buckets"]:
                    data.append([P(b["name"], "fs_ind"), P(money(b["income"]), "cellr"), P(money(b["expense"]), "cellr"),
                                 P(money(b["net"]), "cellr")])
        if other["income"] != "0.00" or other["expense"] != "0.00":
            data.append([P(other["name"], "cell"), P(money(other["income"]), "cellr"), P(money(other["expense"]), "cellr"),
                         P(money(other["net"]), "cellr")])
        data.append([_b("Total"), P(money(t['net_revenue']), "fs_rb"), P(money(t['expenses']), "fs_rb"),
                     P(money(t['net_proceeds']), "fs_rb")])
        f += [Spacer(1, 8), KeepTogether([PM("Net by category / bucket", "h2"),
                                          _table(data, [W - 3.6 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch], bold_rows=bold)])]
    elif s["plain_lines"]:
        data = [_h(["Date", "Contributor / description", "Income", "Expenses"])]
        for x in s["plain_lines"]:
            data.append([P(_us(x["date"]), "cell"), P(x["who"], "cell"), P(money(x["income"]) if x["income"] else "", "cellr"),
                         P(money(x["expense"]) if x["expense"] else "", "cellr")])
        data.append([_b("Total"), "", P(money(t['net_revenue']), "fs_rb"), P(money(t['expenses']), "fs_rb")])
        f += [Spacer(1, 8), PM("Lines", "h2"), _table(data, [0.9 * inch, W - 3.3 * inch, 1.2 * inch, 1.2 * inch])]
    if s["takeaways"]:
        f += [Spacer(1, 8), PM("Key takeaways", "h2")] + [P(f"•  {x}", "body") for x in s["takeaways"]]
    if s["summary_notes"]:
        f += [Spacer(1, 8), PM("Summary notes", "h2"), P(s["summary_notes"], "body")]
    f += [Spacer(1, 10), P("Prepared for committee & stakeholders. Transaction-level evidence (every line and its "
                           "documents) is in the Audit report.", "small")]
    doc.build(f)
    return buf.getvalue(), doc


def _itemized(title: str, lines: list[dict], subtotal: str, W: float) -> list:
    out = [P(title, "metak")]
    if len(lines) > TWO_COLUMN_AFTER:   # long lists: two columns side by side (date, contributor and note, amount)
        half = (len(lines) + 1) // 2
        cw = (W - 12) / 2

        def col(part):
            d = [_h(["Date", "Contributor", "Amount"])]
            for x in part:
                who = escape(x["who"]) + (f"<br/><font size=6>{escape(x['note'])}</font>" if x["note"] else "")
                d.append([P(_us(x["date"]), "cell"), PM(who, "cell"), P(money(x["amount"]), "cellr")])
            return _table(d, [0.75 * inch, cw - 1.75 * inch, 1.0 * inch])
        pair = Table([[col(lines[:half]), col(lines[half:])]], colWidths=[cw + 6, cw + 6])
        pair.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                  ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
        out.append(pair)
    else:
        d = [_h(["Date", "Contributor", "Amount", "Note"])]
        for x in lines:
            d.append([P(_us(x["date"]), "cell"), P(x["who"], "cell"), P(money(x["amount"]), "cellr"), P(x["note"] or "", "cell")])
            d += _charge_rows(x, 4, 2)
        out.append(_table(d, [0.8 * inch, 2.3 * inch, 1.0 * inch, W - 4.1 * inch]))
    out += [P(f"Subtotal: {money(subtotal)}", "cellr"), Spacer(1, 4)]
    return out


def pdf_file(db: Session, ctx, s: dict) -> tuple[str, str, int]:
    pdf, doc = build_pdf(s)
    title = (f"Prepared for committee & stakeholders  ·  Generated {_us(s['generated_at'])}  ·  {s['status']['text']}")
    fd, path = tempfile.mkstemp(prefix="fmpoc-fr-summary-", suffix=".pdf")
    os.close(fd)
    try:
        pages = _stamp_and_write(pdf, doc, path, title)
    except Exception:
        os.unlink(path)
        raise
    return path, f"fundraiser-{safe_name(s['fundraiser'])}-report.pdf", pages


# ------------------------------------------------------------------ Word (.docx)
_HEX_BAND, _HEX_TILE, _HEX_HEAD, _HEX_GROUP = "1F3A5F", "EEF3F9", "DFE7F1", "F2F5F9"


def _shade(cell, hex_color: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def _field(paragraph, code: str) -> None:
    """A PAGE / NUMPAGES field (Word, Google Docs and Pages update it)."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == code:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
            run._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def _cell(cell, text: str, bold=False, right=False, size=9, color=None, indent=0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.left_indent = Inches(0.15 * indent)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def _dtable(d, header: list[str], widths: list[float], right_cols: set[int]):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        _cell(t.rows[0].cells[i], h, bold=True, right=i in right_cols)
        _shade(t.rows[0].cells[i], _HEX_HEAD)
    for i, w in enumerate(widths):
        for c in t.columns[i].cells:
            c.width = Inches(w)
    return t


def _drow(t, values: list[str], right_cols: set[int], bold=False, shade=None, indent=0, widths=None):
    cells = t.add_row().cells
    for i, v in enumerate(values):
        _cell(cells[i], v, bold=bold, right=i in right_cols, indent=indent if i == 0 else 0)
        if shade:
            _shade(cells[i], shade)
        if widths:
            cells[i].width = Inches(widths[i])


def build_docx(s: dict) -> bytes:
    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.7))
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    t, st = s["totals"], s["status"]
    # title band
    band = d.add_table(rows=1, cols=1)
    c = band.rows[0].cells[0]
    _shade(c, _HEX_BAND)
    c.text = ""
    for i, (txt, size, bold) in enumerate(((s["organization"], 10, False), (s["fundraiser"], 18, True),
                                           ("Financial & Proceeds Statement", 11, True),
                                           (f"{event_dates(s)}  ·  {st['text']}", 10, False))):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(txt)
        r.bold, r.font.size, r.font.color.rgb = bold, Pt(size), RGBColor(0xFF, 0xFF, 0xFF)
    if st["code"] == "CANCELLED":
        p = d.add_paragraph()
        r = p.add_run(f"Cancelled: {st['cancel_reason'] or 'no reason given'}. The figures below are what was received "
                      "and spent.")
        r.bold = True
        r.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)
    d.add_heading("Executive summary", level=2)
    tiles = d.add_table(rows=1, cols=3)
    tiles.style = "Table Grid"
    for cell, (k, v, cap) in zip(tiles.rows[0].cells, (
            ("Total net revenue", t["net_revenue"], "All income counted for the fundraiser"
             + ("; cash float returned excluded" if t["cash_float_returned"] != "0.00" else "")),
            ("Total expenses", t["expenses"], "Fundraiser expenses" + ("; cash float not included" if t["cash_float_out"] != "0.00" else "")),
            ("Net proceeds", t["net_proceeds"], f"Net margin {_p(t['net_margin'])} of net revenue"))):
        _shade(cell, _HEX_TILE)
        _cell(cell, k, bold=True, size=9)
        p = cell.add_paragraph()
        r = p.add_run(money(v))
        r.bold, r.font.size = True, Pt(16)
        p2 = cell.add_paragraph()
        r2 = p2.add_run(cap)
        r2.font.size = Pt(8)
    d.add_heading("Financial summary", level=2)
    W3 = [1.8, 3.6, 1.3]
    tb = _dtable(d, ["Component", "Details", "Amount"], W3, {2})
    _drow(tb, ["Gross deposits", "All deposits counted for the fundraiser", money(t["gross_deposits"])], {2}, widths=W3)
    if t["cash_float_returned"] != "0.00":
        _drow(tb, ["Less cash float returned", "Cash float coming back inside deposits; not income",
                   f"-{money(t['cash_float_returned'])}"], {2}, widths=W3)
    _drow(tb, ["Net revenue", "", money(t["net_revenue"])], {2}, bold=True, widths=W3)
    _drow(tb, ["Less expenses", "Fundraiser expenses" + ("; cash float taken out is not an expense"
                                                         if t["cash_float_out"] != "0.00" else ""),
               f"-{money(t['expenses'])}"], {2}, widths=W3)
    _drow(tb, ["Net proceeds", f"Net margin {_p(t['net_margin'])}", money(t["net_proceeds"])], {2}, bold=True,
          shade=_HEX_GROUP, widths=W3)
    p = d.add_paragraph(f"ROI (net proceeds ÷ expenses): {t['roi'] if t['roi'] is not None else 'n/a (no expenses)'}")
    p.runs[0].font.size = Pt(8)
    for y in s["per_fiscal_year"]:
        q = d.add_paragraph(f"{y['fiscal_year']}: revenue {money(y['income'])}, expenses {money(y['expense'])}, net {money(y['net'])}")
        q.runs[0].font.size = Pt(8)
    groups, other = s["groups"], s["other"]
    if s["has_buckets"]:
        d.add_heading("Revenue breakdown", level=2)
        W4 = [3.5, 0.8, 1.2, 1.2]
        tb = _dtable(d, ["Category / bucket", "Lines", "Subtotal", "% of revenue"], W4, {1, 2, 3})
        for g in groups:
            if g["income"] == "0.00":
                continue
            _drow(tb, [g["name"], str(g["income_lines"]), money(g["income"]), _p(g["income_pct"])], {1, 2, 3}, bold=True,
                  shade=_HEX_GROUP, widths=W4)
            if g["kind"] == "category":
                for b in g["buckets"]:
                    if b["income"] != "0.00":
                        _drow(tb, [b["name"], str(len(b["income_lines"])), money(b["income"]), _p(b["income_pct"])],
                              {1, 2, 3}, indent=1, widths=W4)
        if other["income"] != "0.00":
            _drow(tb, [other["name"], str(len(other["income_lines"])), money(other["income"]), _p(other["income_pct"])],
                  {1, 2, 3}, widths=W4)
        _drow(tb, ["Total net revenue", "", money(t["net_revenue"]), "100.0%" if t["net_revenue"] != "0.00" else "n/a"],
              {1, 2, 3}, bold=True, widths=W4)
        if s["options"]["itemized"]:
            d.add_heading("Itemized revenue", level=2)
            subs = [((b["name"] if g["kind"] == "bucket" else f"{g['name']} — {b['name']}"), b["income_lines"], b["income"])
                    for g in groups for b in g["buckets"] if b["income_lines"]]
            if other["income_lines"]:
                subs.append((other["name"], other["income_lines"], other["income"]))
            WI = [0.9, 2.4, 1.1, 2.3]
            for title, lines, sub in subs:
                q = d.add_paragraph()
                q.add_run(title).bold = True
                tb = _dtable(d, ["Date", "Contributor", "Amount", "Note"], WI, {2})
                for x in lines:
                    _drow(tb, [_us(x["date"]), x["who"], money(x["amount"]), x["note"] or ""], {2}, widths=WI)
                    if x.get("less"):
                        _drow(tb, ["Charge", "", money(x["charge"]), ""], {2}, indent=2, widths=WI)
                        _drow(tb, [f"Less: {x['less_description'] or ''}", "", f"-{money(x['less'])}", ""], {2}, indent=2, widths=WI)
                _drow(tb, ["Subtotal", "", money(sub), ""], {2}, bold=True, widths=WI)
        d.add_heading("Expense breakdown", level=2)
        WE = [4.2, 1.25, 1.25]
        tb = _dtable(d, ["Category / bucket / line", "Amount", "% of expenses"], WE, {1, 2})
        for g in groups:
            if g["expense"] == "0.00":
                continue
            _drow(tb, [g["name"], money(g["expense"]), _p(g["expense_pct"])], {1, 2}, bold=True, shade=_HEX_GROUP, widths=WE)
            for b in g["buckets"]:
                if not b["expense_lines"]:
                    continue
                if g["kind"] == "category":
                    _drow(tb, [b["name"], money(b["expense"]), _p(b["expense_pct"])], {1, 2}, indent=1, widths=WE)
                for x in b["expense_lines"]:
                    for label, amt in expense_line_rows(x):
                        _drow(tb, [label, amt, ""], {1, 2}, indent=2, widths=WE)
        if other["expense_lines"]:
            _drow(tb, [other["name"], money(other["expense"]), _p(other["expense_pct"])], {1, 2}, bold=True, shade=_HEX_GROUP, widths=WE)
            for x in other["expense_lines"]:
                _drow(tb, [x["who"], money(x["amount"]), ""], {1, 2}, indent=2, widths=WE)
        _drow(tb, ["Total expenses", money(t["expenses"]), "100.0%" if t["expenses"] != "0.00" else "n/a"], {1, 2}, bold=True, widths=WE)
        if t["cash_float_out"] != "0.00":
            q = d.add_paragraph(f"Cash float taken out ({money(t['cash_float_out'])}) is not an expense; "
                                f"{money(t['cash_float_returned'])} came back in the deposits.")
            q.runs[0].font.size = Pt(8)
        d.add_heading("Net by category / bucket", level=2)
        WN = [3.1, 1.2, 1.2, 1.2]
        tb = _dtable(d, ["Category / bucket", "Income", "Expenses", "Net"], WN, {1, 2, 3})
        for g in groups:
            _drow(tb, [g["name"], money(g["income"]), money(g["expense"]), money(g["net"])], {1, 2, 3}, bold=True,
                  shade=_HEX_GROUP, widths=WN)
            if g["kind"] == "category":
                for b in g["buckets"]:
                    _drow(tb, [b["name"], money(b["income"]), money(b["expense"]), money(b["net"])], {1, 2, 3}, indent=1, widths=WN)
        if other["income"] != "0.00" or other["expense"] != "0.00":
            _drow(tb, [other["name"], money(other["income"]), money(other["expense"]), money(other["net"])], {1, 2, 3}, widths=WN)
        _drow(tb, ["Total", money(t["net_revenue"]), money(t["expenses"]), money(t["net_proceeds"])], {1, 2, 3}, bold=True, widths=WN)
    elif s["plain_lines"]:
        d.add_heading("Lines", level=2)
        WP = [0.9, 3.4, 1.2, 1.2]
        tb = _dtable(d, ["Date", "Contributor / description", "Income", "Expenses"], WP, {2, 3})
        for x in s["plain_lines"]:
            _drow(tb, [_us(x["date"]), x["who"], money(x["income"]) if x["income"] else "",
                       money(x["expense"]) if x["expense"] else ""], {2, 3}, widths=WP)
        _drow(tb, ["Total", "", money(t["net_revenue"]), money(t["expenses"])], {2, 3}, bold=True, widths=WP)
    if s["takeaways"]:
        d.add_heading("Key takeaways", level=2)
        for x in s["takeaways"]:
            d.add_paragraph(x, style="List Bullet")
    if s["summary_notes"]:
        d.add_heading("Summary notes", level=2)
        for para in s["summary_notes"].split("\n"):
            d.add_paragraph(para)
    q = d.add_paragraph("Prepared for committee & stakeholders. Transaction-level evidence (every line and its "
                        "documents) is in the Audit report.")
    q.runs[0].font.size = Pt(8)
    # headings in the report's colour
    for name in ("Heading 2",):
        hs = d.styles[name]
        hs.font.name = "Arial"
        hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(f"Prepared for committee & stakeholders  ·  Generated {_us(s['generated_at'])}  ·  {st['text']}  ·  Page ")
    r.font.size = Pt(7)
    _field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(7)
    _field(fp, "NUMPAGES")
    d.core_properties.title = f"{s['fundraiser']} — Financial & Proceeds Statement"
    d.core_properties.author = s["organization"]
    d.core_properties.comments = ""
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def docx_name(s: dict) -> str:
    return f"fundraiser-{safe_name(s['fundraiser'])}-report.docx"


__all__ = ["summary", "report_status", "build_pdf", "pdf_file", "build_docx", "docx_name", "_FONT"]
