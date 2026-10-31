"""2.1.0 (#277): the Fiscal Year fundraising report (PDF) and presentation (.pptx), drawn from
`fundraiser_series.fy_report()`. The PDF and the deck show the same content; the deck adds charts. The deck follows
the #276 compatibility rules (Arial, native charts with data labels, speaker notes, no pictures or media)."""
from __future__ import annotations

import io
import os
import tempfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from . import fundraiser_deck as deck
from . import branding
from .fundraiser_summary import _b, _band, _tile, _h, _p, _table, _us, heading, safe_name
from .reports import FRAME_W, P, S, _AuditDoc, _Mark, _stamp_and_write, money

NO_HISTORY = "No earlier years linked yet."


def _growth(v: str | None) -> str:
    if v is None:
        return "n/a"
    return f"{'+' if not v.startswith('-') else ''}{v}%"


def _chg(x: dict) -> str:
    if x["change"] is None:
        return "—"
    if x["change"] in ("0.00", "-0.00"):
        return "no change"
    word = "up" if not x["change"].startswith("-") else "down"
    return f"{word} {money(x['change'].lstrip('-'))}" + (f" ({x['change_pct'].lstrip('-')}%)" if x.get("change_pct") else "")


def highlights(rep: dict) -> list[tuple[str, str, str]]:
    out = [("Raised this Fiscal Year (net)", money(rep["year_total"]),
            ("YTD — the year is still open" if not rep["fy_ended"] else "Net proceeds of the fundraisers shown"))]
    if rep["linked_history"]:
        out.append(("Growth vs last year", _growth(rep["growth_vs_last"]),
                    f"and {_growth(rep['growth_vs_first'])} vs {rep['period'][0]}"))
    if rep["top"]:
        out.append(("Top fundraiser", rep["top"]["name"], f"{money(rep['top']['net'])} · {_p(rep['top']['share'])} of the year"))
    out.append((f"Cumulative {rep['period'][0]}–{rep['year']}", money(rep["cumulative"]), f"Peak year {rep['peak_year']}"))
    return out


def est_lines(rep: dict) -> list[str]:
    """2.3.0 (#390): "* Estimated: <year> <series> — <note>" lines (also in the report's footnotes)."""
    return [x for x in rep["footnotes"] if x.startswith("* Estimated:")]


def _star(est: bool) -> str:
    return "*" if est else ""


def file_stem(org: str, rep: dict) -> str:
    return f"fundraising-{safe_name(rep['fiscal_year']['display_name'])}"


# ------------------------------------------------------------------ PDF
def yoy_layout(W: float, years: list, cells: list[str]) -> tuple[list, list[float]]:
    """How many of the latest `years` (up to 6) "Year over year by series" shows, and its column widths (Series, one per
    year, change), so that it always fits the frame and no amount wraps: a year column is as wide as the widest amount
    in the table plus padding (at most 0.95"), the change column 1.15" and the Series column at least 1.1"; years that
    don't fit are left out from the oldest. Before #295 the years were 0.95" each regardless, so 6 years left the
    Series column a negative width and the PDF failed."""
    from . import reports
    change, series_min, pad = 1.15 * inch, 1.1 * inch, 10
    need = max([stringWidth(c, reports._FONT, 7.5) for c in cells] + [stringWidth("0000", reports._FONT_BOLD, 7.5)]) + pad
    years = list(years)[-6:]
    fit = max(1, int((W - change - series_min) // need))
    years = years[-fit:]
    yr = min(0.95 * inch, max(need, (W - change - series_min) / len(years)))
    return years, [W - change - yr * len(years)] + [yr] * len(years) + [change]


def build_pdf(org: str, rep: dict) -> tuple[bytes, _AuditDoc]:
    buf = io.BytesIO()
    doc = _AuditDoc(buf, title="Fundraising report", author=org)
    W = FRAME_W - 12
    band = Table([[[P(org, "fs_w"), Paragraph(escape(f"Fundraising — {rep['fiscal_year']['display_name']}"),
                                             S["fs_wt"]),
                    P("Fiscal Year fundraising report", "fs_wb"),
                    P(f"History {rep['period'][0]}–{rep['year']}  ·  cumulative {money(rep['cumulative'])}  ·  "
                      f"peak year {rep['peak_year']}", "fs_w")]]], colWidths=[W])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _band()), ("LEFTPADDING", (0, 0), (-1, -1), 12),
                              ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    f: list = [_Mark(doc, rep["fiscal_year"]["display_name"]), band, Spacer(1, 8), *heading("Executive highlights")]
    hl = highlights(rep)
    cells = [[[P(k, "fs_tk"), P(v, "fs_tv"), P(c, "small")] for k, v, c in hl]]
    tt = Table(cells, colWidths=[W / len(hl)] * len(hl))
    tt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _tile()), ("BOX", (0, 0), (-1, -1), 0.5, branding.color("outline_variant")),
                            ("INNERGRID", (0, 0), (-1, -1), 0.5, branding.color("outline_variant")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    f += [tt, Spacer(1, 8), *heading("This year by fundraiser")]
    data = [_h(["Fundraiser", "Event", "Status", "Net revenue", "Expenses", "Net", "Margin", "Share"])]
    for r in rep["this_year"]:
        ev = _us(r["start_date"]) if r["start_date"] == r["end_date"] else f"{_us(r['start_date'])}–{_us(r['end_date'])}"
        data.append([P(r["name"], "cell"), P(ev, "cell"), P(r["status"], "cell"), P(money(r["net_revenue"]), "cellr"),
                     P(money(r["expenses"]), "cellr"), P(money(r["net_proceeds"]), "cellr"), P(_p(r["net_margin"]), "cellr"),
                     P(_p(r["share"]), "cellr")])
    data.append([_b("Total"), "", "", "", "", P(money(rep["year_total"]), "fs_rb"), "", P("100.0%", "fs_rb")])
    f.append(_table(data, [W - 5.9 * inch, 1.1 * inch, 0.9 * inch, 0.85 * inch, 0.8 * inch, 0.85 * inch, 0.65 * inch, 0.75 * inch],
                    bold_rows=[len(data) - 1]))
    if rep.get("events"):           # 2.4.0 (#429): Events, separate from the fundraising figures
        f += [Spacer(1, 8), *heading("Events (not fundraisers)")]
        data = [_h(["Event", "Dates", "Status", "Income", "Expenses", "Net"])]
        for r in rep["events"]:
            ev = _us(r["start_date"]) if r["start_date"] == r["end_date"] else f"{_us(r['start_date'])}–{_us(r['end_date'])}"
            data.append([P(r["name"], "cell"), P(ev, "cell"), P(r["status"], "cell"), P(money(r["income"]), "cellr"),
                         P(money(r["expenses"]), "cellr"), P(money(r["net"]), "cellr")])
        data.append([_b("Total"), "", "", "", "", P(money(rep["events_net"]), "fs_rb")])
        f.append(_table(data, [W - 4.6 * inch, 1.1 * inch, 0.9 * inch, 0.85 * inch, 0.85 * inch, 0.9 * inch],
                        bold_rows=[len(data) - 1]))
        f.append(P("Events are not counted in the fundraising totals, shares or history above.", "small"))
    # annual trend (table; the deck has the chart)
    f += [Spacer(1, 8), *heading("Annual trend")]
    if rep["linked_history"]:
        data = [_h(["Year", "Net raised"])]
        for t in rep["trend"]:
            label = str(t["year"]) + _star(t.get("estimated")) + (" (YTD)" if t["current"] and not rep["fy_ended"] else "")
            data.append([P(label, "cell"), P(money(t["total"]), "cellr")])
        if rep["targeted"] and rep["estimates"]:
            data.append([P(f"{rep['year']} targeted (YTD + estimate {money(rep['estimate_total'])})", "cell"),
                         P(money(rep["targeted_total"]), "fs_rb")])
        f.append(_table(data, [W - 1.5 * inch, 1.5 * inch]))
        f += [P(x, "small") for x in est_lines(rep)]
    else:
        f.append(P(NO_HISTORY, "body"))
    if rep["targeted"] and rep["estimates"]:
        f.append(P("Estimates: " + "; ".join(f"{e['series']} {money(e['estimate'])} (average of {e['years']} year"
                                              f"{'s' if e['years'] != 1 else ''} {money(e['average'])})" for e in rep["estimates"]), "small"))
    # share by series
    f += [Spacer(1, 8), *heading("Share by series over the period")]
    if rep["share"]:
        data = [_h(["Series", "Total", "Share"])]
        for x in rep["share"]:
            name = x["name"] + (f" ({', '.join(x['members'])})" if x.get("members") else "")
            data.append([P(name, "cell"), P(money(x["total"]), "cellr"), P(_p(x["share"]), "cellr")])
        f.append(_table(data, [W - 2.4 * inch, 1.2 * inch, 1.2 * inch]))
    # average per active year
    f += [Spacer(1, 8), *heading("Average per active year")]
    if rep["linked_history"]:
        data = [_h(["Series", "Total over the period", "Active years", "Average per active year"])]
        for a in rep["averages"]:
            data.append([P(a["name"], "cell"), P(money(a["total"]), "cellr"), P(str(a["active_years"]), "cellr"),
                         P(money(a["average"]), "cellr")])
        f.append(_table(data, [W - 4.2 * inch, 1.5 * inch, 1.0 * inch, 1.7 * inch]))
    else:
        f.append(P(NO_HISTORY, "body"))
    # year over year by series (#277 feedback)
    if rep["year_over_year"]:
        f += [Spacer(1, 8), *heading("Year over year by series"),
              P("Net raised by each event in each year, and the change on the year before — to help decide which "
                "fundraisers to keep offering.", "small"), Spacer(1, 3)]
        cells = [money(v) for x in rep["year_over_year"] for v in x["values"] if v is not None]
        years, widths = yoy_layout(W, rep["period"], cells)
        idx = [rep["period"].index(y) for y in years]
        data = [_h(["Series"] + [str(y) for y in years] + ["Change vs year before"])]
        for x in rep["year_over_year"]:
            data.append([P(x["name"], "cell")] + [P(money(x["values"][i]) + _star(x.get("estimated", [False] * 99)[i])
                                                    if x["values"][i] is not None else "—", "cellr") for i in idx]
                        + [P(_chg(x), "cell")])
        f.append(_table(data, widths))
        f += [P(x, "small") for x in est_lines(rep)]
        if len(rep["period"]) > len(years):
            f.append(P(f"The last {len(years)} years are shown; the historical results below list every year.", "small"))
    # historical table
    f += [Spacer(1, 8), *heading("Historical results")]
    data = [_h(["Year", "Series", "Amount raised", "Source", "Note"])]
    for h in rep["history"]:
        data.append([P(str(h["year"]) + _star(h.get("estimated")), "cell"), P(h["series"], "cell"), P(money(h["amount"]), "cellr"),
                     P(h["source"], "cell"), P(h["note"] or "", "cell")])
    f.append(_table(data, [0.6 * inch, 2.3 * inch, 1.2 * inch, 1.2 * inch, W - 5.3 * inch]))
    f += [P(x, "small") for x in est_lines(rep)]
    if rep["notes"]:
        f += [Spacer(1, 8), *heading("Notes"), P(rep["notes"], "body")]
    if rep["footnotes"]:
        f += [Spacer(1, 8)] + [P(x, "small") for x in rep["footnotes"]]
    doc.build(f)
    return buf.getvalue(), doc


def pdf_file(org: str, rep: dict) -> tuple[str, str, int]:
    pdf, doc = build_pdf(org, rep)
    title = f"{org}  ·  Fiscal Year fundraising report  ·  Generated {_us(rep['generated_at'])}"
    fd, path = tempfile.mkstemp(prefix="fmpoc-fr-fy-", suffix=".pdf")
    os.close(fd)
    try:
        pages = _stamp_and_write(pdf, doc, path, title)
    except Exception:
        os.unlink(path)
        raise
    return path, file_stem(org, rep) + "-report.pdf", pages


# ------------------------------------------------------------------ presentation
def build_pptx(org: str, rep: dict) -> bytes:
    prs = Presentation()
    deck._strip_template_extras(prs)
    prs.slide_width, prs.slide_height = deck.W, deck.H
    blank = prs.slide_layouts[6]
    fyn = rep["fiscal_year"]["display_name"]
    # 1. title
    sl = prs.slides.add_slide(blank)
    deck._mark(sl, Inches(0.6), Inches(1.6), Inches(1.3), Inches(1.7))
    deck._text(sl, Inches(2.4), Inches(1.45), Inches(10), Inches(0.7), org, size=28, bold=True)
    deck._text(sl, Inches(2.4), Inches(2.2), Inches(10), Inches(1.2), f"Fundraising — {fyn}", size=40, bold=True, color=deck.C.BLUE)
    deck._text(sl, Inches(2.4), Inches(3.5), Inches(10), Inches(0.6),
               f"History {rep['period'][0]}–{rep['year']} · cumulative {money(rep['cumulative'])} · peak year {rep['peak_year']}",
               size=18, color=deck.C.GREY)
    deck._notes(sl, [f"Fiscal Year fundraising report for {fyn}, {len(rep['this_year'])} fundraisers selected.",
                     f"Cumulative net raised {rep['period'][0]}–{rep['year']}: {money(rep['cumulative'])}; peak year {rep['peak_year']} "
                     f"({money(rep['peak_total'])})."])
    # 2. highlights
    sl = prs.slides.add_slide(blank)
    deck._title(sl, "Executive highlights")
    hl = highlights(rep)
    w = 12.1 / len(hl)
    for i, (k, v, c) in enumerate(hl):
        x = Inches(0.6 + i * w)
        card = sl.shapes.add_shape(1, x, Inches(2.0), Inches(w - 0.25), Inches(2.6))  # rectangle
        card.fill.solid()
        card.fill.fore_color.rgb = deck.C.LIGHT
        card.line.color.rgb = deck.C.LINE
        deck._text(sl, x + Inches(0.2), Inches(2.15), Inches(w - 0.6), Inches(0.6), k, size=15, bold=True, color=deck.C.GREY)
        deck._text(sl, x + Inches(0.2), Inches(2.8), Inches(w - 0.6), Inches(1.0), v, size=26, bold=True)
        deck._text(sl, x + Inches(0.2), Inches(3.85), Inches(w - 0.6), Inches(0.6), c, size=12, color=deck.C.GREY)
    deck._notes(sl, [f"Raised this Fiscal Year = sum of the selected fundraisers' net proceeds = {money(rep['year_total'])}.",
                     f"Growth vs last year = (this year − last year) ÷ last year = {_growth(rep['growth_vs_last'])}.",
                     f"Growth vs {rep['period'][0]} = {_growth(rep['growth_vs_first'])}."]
               + ([f"Top fundraiser: {rep['top']['name']} {money(rep['top']['net'])} = {_p(rep['top']['share'])} of the year."]
                  if rep["top"] else []))
    # 3. this year by fundraiser (table, 12 rows per slide)
    rows = [[r["name"], money(r["net_revenue"]), money(r["expenses"]), money(r["net_proceeds"]), _p(r["net_margin"]), _p(r["share"])]
            for r in rep["this_year"]] + [["Total", "", "", money(rep["year_total"]), "", "100.0%"]]
    for k, part in enumerate(deck._chunks(rows, deck.TABLE_ROWS)):
        sl = prs.slides.add_slide(blank)
        deck._title(sl, f"{fyn} by fundraiser" + (" (continued)" if k else ""))
        deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [["Fundraiser", "Net revenue", "Expenses", "Net", "Margin", "Share"]] + part,
                    [5.1, 1.6, 1.4, 1.6, 1.1, 1.3], {1, 2, 3, 4, 5},
                    bold_rows=[i + 1 for i, r in enumerate(part) if r[0] == "Total"])
        deck._notes(sl, ["Each row is that fundraiser's stakeholder report: net proceeds = net revenue − expenses; margin = "
                         "net ÷ net revenue; share = its net ÷ the year's total."])
    # 3b. 2.4.0 (#429): Events, a section of their own
    if rep.get("events"):
        rows = [[r["name"], money(r["income"]), money(r["expenses"]), money(r["net"])] for r in rep["events"]]
        rows.append(["Total", "", "", money(rep["events_net"])])
        for k, part in enumerate(deck._chunks(rows, deck.TABLE_ROWS)):
            sl = prs.slides.add_slide(blank)
            deck._title(sl, f"{fyn} events" + (" (continued)" if k else ""), "Not counted in the fundraising figures")
            deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [["Event", "Income", "Expenses", "Net"]] + part,
                        [6.1, 2.0, 2.0, 2.0], {1, 2, 3}, bold_rows=[i + 1 for i, r in enumerate(part) if r[0] == "Total"])
            deck._notes(sl, ["Events (not fundraisers): income, expenses and net for each; they are listed separately "
                             "and do not count in the year's fundraising total."])
    # 4. annual trend
    sl = prs.slides.add_slide(blank)
    deck._title(sl, "Annual trend", "Net raised per year")
    if rep["linked_history"]:
        cd = CategoryChartData(number_format='"$"#,##0.00')
        cats = [str(t["year"]) + _star(t.get("estimated")) + (" (YTD)" if t["current"] and not rep["fy_ended"] else "")
                for t in rep["trend"]]
        cd.categories = cats
        cd.add_series("Recorded / entered", [float(t["total"]) for t in rep["trend"]])
        est = rep["targeted"] and rep["estimates"]
        if est:
            cd.add_series("Estimate (targeted)", [float(rep["estimate_total"]) if t["current"] else 0 for t in rep["trend"]])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED if est else XL_CHART_TYPE.COLUMN_CLUSTERED,
                                 Inches(0.9), Inches(1.8), Inches(11.5), Inches(5.2), cd).chart
        deck._style_chart(ch, legend=bool(est))
        deck._labels(ch.plots[0], pos=None if est else XL_LABEL_POSITION.OUTSIDE_END, size=12)
        ch.plots[0].series[0].format.fill.solid()
        ch.plots[0].series[0].format.fill.fore_color.rgb = deck.C.SERIES[0]
        if est:
            ch.plots[0].series[1].format.fill.solid()
            ch.plots[0].series[1].format.fill.fore_color.rgb = deck.C.SERIES[1]
        for i, t in enumerate(rep["trend"]):     # 2.3.0 (#390): an estimated year's bar lighter, labelled *
            if t.get("estimated"):
                pt = ch.plots[0].series[0].points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = deck._lighter(deck.C.SERIES[0])
        if est_lines(rep):
            deck._text(sl, Inches(0.9), Inches(7.0), Inches(11.5), Inches(0.45), est_lines(rep), size=10, color=deck.C.GREY)
        deck._notes(sl, est_lines(rep) + [f"{t['year']}: {money(t['total'])}" for t in rep["trend"]]
                    + ([f"Targeted total {money(rep['targeted_total'])} = YTD + estimate {money(rep['estimate_total'])}: "
                        + "; ".join(f"{e['series']} average of {e['years']} years" for e in rep["estimates"])] if est else []))
    else:
        deck._text(sl, Inches(0.9), Inches(2.4), Inches(11), Inches(1), NO_HISTORY, size=22, color=deck.C.GREY)
        deck._notes(sl, ["Link fundraisers with “Continues from”, or enter earlier years in the Fundraiser series list."])
    # 5. share by series
    if rep["share"]:
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Share by series", f"{rep['period'][0]}–{rep['year']}")
        cd = CategoryChartData(number_format='"$"#,##0.00')
        cd.categories = [f"{x['name']} — {money(x['total'])} ({_p(x['share'])})" for x in rep["share"]]
        cd.add_series("Net raised", [max(float(x["total"]), 0) for x in rep["share"]])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.6), cd).chart
        deck._style_chart(ch, legend_pos=XL_LEGEND_POSITION.RIGHT)
        deck._labels(ch.plots[0], fmt="0.0%", size=12)
        ch.plots[0].data_labels.show_percentage = True
        ch.plots[0].data_labels.show_value = False
        for i, pt in enumerate(ch.plots[0].series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = deck.C.SERIES[i % len(deck.C.SERIES)]
        deck._notes(sl, [f"Share of the period total {money(rep['cumulative'])}:"]
                    + [f"{x['name']}: {money(x['total'])} ({_p(x['share'])})" + (f" — {', '.join(x['members'])}" if x.get("members") else "")
                       for x in rep["share"]])
    # 6. average per active year
    sl = prs.slides.add_slide(blank)
    deck._title(sl, "Average per active year")
    if rep["linked_history"] and rep["averages"]:
        items = rep["averages"][:deck.TABLE_ROWS]
        cd = CategoryChartData(number_format='"$"#,##0.00')
        cd.categories = [a["name"] for a in items]
        cd.add_series("Average per active year", [float(a["average"]) for a in items])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.6), Inches(6.9), Inches(5.6), cd).chart
        deck._style_chart(ch, legend=False)
        deck._labels(ch.plots[0], pos=XL_LABEL_POSITION.OUTSIDE_END, size=11)
        rows = [["Series", "Years", "Average"]] + [[a["name"], str(a["active_years"]), money(a["average"])] for a in items]
        deck._table(sl, Inches(7.6), Inches(1.7), Inches(5.3), rows, [3.0, 0.9, 1.4], {1, 2}, size=11)
        deck._notes(sl, ["Average per active year = total over the period ÷ the years the series was active."]
                    + [f"{a['name']}: {money(a['total'])} over {a['active_years']} years = {money(a['average'])}" for a in rep["averages"]])
    else:
        deck._text(sl, Inches(0.9), Inches(2.4), Inches(11), Inches(1), NO_HISTORY, size=22, color=deck.C.GREY)
        deck._notes(sl, [NO_HISTORY])
    # year over year by series (#277 feedback): a line per series, and the change table
    if rep["year_over_year"]:
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Year over year by series", "Net raised per year")
        cd = CategoryChartData(number_format='"$"#,##0.00')
        est_years = {y for t in rep["trend"] if t.get("estimated") for y in [t["year"]]}
        cd.categories = [str(y) + _star(y in est_years) for y in rep["period"]]
        for x in rep["year_over_year"][:8]:
            cd.add_series(x["name"], [float(v) if v is not None else None for v in x["values"]])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), cd).chart
        deck._style_chart(ch)
        deck._labels(ch.plots[0], pos=XL_LABEL_POSITION.ABOVE, size=9)
        for i, ser in enumerate(ch.plots[0].series):
            ser.format.line.color.rgb = deck.C.SERIES[i % len(deck.C.SERIES)]
            ser.smooth = False
        deck._notes(sl, ["Net raised by each series per year; a gap = not held that year."] + est_lines(rep)
                    + [f"{x['name']}: " + ", ".join(f"{y} {money(v) if v is not None else '—'}" for y, v in zip(rep["period"], x["values"]))
                       + f"; vs year before: {_chg(x)}" for x in rep["year_over_year"]])
        rows = [["Series", f"{rep['year']}", "Change vs year before"]] + [
            [x["name"], money(x["values"][-1]) if x["values"][-1] is not None else "—", _chg(x)] for x in rep["year_over_year"]]
        for k, part in enumerate(deck._chunks(rows[1:], deck.TABLE_ROWS)):
            sl = prs.slides.add_slide(blank)
            deck._title(sl, "Change on the year before" + (" (continued)" if k else ""))
            deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [rows[0]] + part, [6.1, 2.5, 3.5], {1}, size=13)
            deck._notes(sl, ["Change = this year's net − the series' most recent earlier year; % of that earlier year."])
    # 7. historical table (about 12 rows per slide)
    rows = [[str(h["year"]) + _star(h.get("estimated")), h["series"], money(h["amount"]), h["source"], h["note"] or ""]
            for h in rep["history"]]
    for k, part in enumerate(deck._chunks(rows, deck.TABLE_ROWS)):
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Historical results" + (" (continued)" if k else ""))
        deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [["Year", "Series", "Amount raised", "Source", "Note"]] + part,
                    [1.0, 4.0, 1.9, 1.8, 3.4], {2}, size=12)
        deck._notes(sl, ["Recorded = from the fundraisers in PennyWarden; Entered = entered by hand for a year before PennyWarden."]
                    + est_lines(rep))
    # 8. notes and footnotes
    lines = ([rep["notes"]] if rep["notes"] else []) + rep["footnotes"]
    if lines:
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Notes")
        deck._text(sl, Inches(0.9), Inches(1.8), Inches(11.5), Inches(5), lines, size=16, color=deck.C.TEXT)
        deck._notes(sl, lines)
    # closing
    sl = prs.slides.add_slide(blank)
    deck._mark(sl, Inches(5.17), Inches(1.7), Inches(1.2), Inches(3.0))
    deck._text(sl, Inches(0.6), Inches(3.2), Inches(12.1), Inches(1.2), "Questions?", size=48, bold=True, align=PP_ALIGN.CENTER)
    deck._text(sl, Inches(0.6), Inches(4.5), Inches(12.1), Inches(0.6), f"{org} · {fyn}", size=20, color=deck.C.GREY, align=PP_ALIGN.CENTER)
    deck._notes(sl, [f"Generated {_us(rep['generated_at'])}."])
    prs.core_properties.title = f"Fundraising — {fyn}"
    prs.core_properties.author = org
    out = io.BytesIO()
    deck.brand_footers(prs)
    prs.save(out)
    return out.getvalue()


__all__ = ["build_pdf", "pdf_file", "build_pptx", "file_stem", "KeepTogether"]
