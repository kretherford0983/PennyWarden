"""2.3.0 (#373): the Quarterly / Annual Fiscal Year financial report as a PDF (tables) and a presentation (.pptx,
charts + tables, commentary in the speaker notes), drawn from `period_report.build()`. Same toolkit as the Fiscal Year
fundraising report: `services/reports.py` for the PDF, `fundraiser_deck` helpers for the deck (#276 rules: Arial,
native charts with data labels, speaker notes, no pictures)."""
from __future__ import annotations

import datetime as dt
import io
import os
import tempfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from . import branding
from . import fundraiser_deck as deck
from .common import GROSS_BUDGET_NOTE
from .fundraiser_summary import _band, _table, _tile, heading, safe_name

from .reports import FRAME_W, P, S, _AuditDoc, _Mark, _stamp_and_write
from .reports import money as cm

DASH = "—"


def _d(iso: str) -> str:
    d = dt.date.fromisoformat(iso)
    return f"{d.strftime('%B')} {d.day}, {d.year}"


def title(rep: dict) -> str:
    fy = rep["fiscal_year"]["display_name"]
    return f"{fy} Annual financial report" if rep["annual"] else f"{fy} {rep['period']} financial report"


def as_of(rep: dict) -> str:
    return (f"Preliminary — as of {_d(rep['cutoff'])}" if rep["preliminary"] else f"As of {_d(rep['cutoff'])}")


def file_stem(org: str, rep: dict) -> str:
    return safe_name(f"{org}-{rep['fiscal_year']['display_name']}-{rep['period_label']}"
                     + ("-preliminary" if rep["preliminary"] else ""))


def _m(c: int | None) -> str:
    return DASH if c is None else cm(c)


def _pct(x: float | None) -> str:
    return DASH if x is None else f"{x:+.1f}%"


def highlight_tiles(rep: dict) -> list[tuple[str, str, str]]:
    h = rep["highlights"]
    per = "the year" if rep["annual"] else rep["period"]
    out = [("Total cash", cm(h["total_cash"]), f"All accounts at {_d(rep['cutoff'])}"),
           ("Net cash change", cm(h["net_cash_change"]), f"Accounts, {per}"),
           ("Budgets over budget", str(h["budgets_over"]), f"{h['budgets_within']} within budget, year to date")]
    if h["fundraisers_on"]:
        out.append(("Fundraiser net", cm(h["fundraiser_net"]),
                    "Fundraisers of the year" if rep["annual"] else f"Fundraisers that ended in {rep['period']}"))
    return out


def notes_lines(rep: dict) -> list[str]:
    """Rules printed under the report (and in the deck's notes)."""
    lines = [f"Everything is as of {_d(rep['cutoff'])}: transactions dated later are not included."]
    if rep["preliminary"]:
        lines.append("Preliminary: the period" + (" or the Fiscal Year" if rep["annual"] else "")
                     + " has not ended" + (" or is not closed" if rep["annual"] else "") + "; figures can still change.")
    if not rep["annual"]:
        lines.append("Earlier quarters of the year are included only in the year-to-date figures.")
    lines.append("Cash flow counts money in (deposits) and money out (withdrawals) of every register account; "
                 "transfers between accounts (Budget 0) are left out.")
    return lines


# ------------------------------------------------------------------ PDF
def _hdr(cells: list[str]) -> list:
    """Header cells: the first left-aligned, the figures right-aligned (as their columns)."""
    return [P(c, "cellb" if i == 0 else "fs_rb") for i, c in enumerate(cells)]


def _grid(data, widths, bold_last=True):
    return _table(data, widths, bold_rows=[len(data) - 1] if bold_last else [])


def _budget_table(rep: dict, key: str, W: float) -> list:
    b = rep["budgets"]
    sec = b[key]
    prior = b["has_prior"]
    if rep["annual"]:
        head = ["Budget", *b["columns"], "Actual", "Budgeted", "Over / under"] + (["Prior year"] if prior else [])
    else:
        head = ["Budget", "Quarter", "Budgeted", "YTD actual", "Over / under"] + (["Prior quarter", "Prior YTD"] if prior else [])
    data = [_hdr(head)]

    def cells(x: dict, label: str) -> list:
        if rep["annual"]:
            row = [P(label, "cell"), *[P(cm(v), "cellr") for v in x["columns"]], P(cm(x["ytd"]), "cellr"),
                   P(cm(x["budgeted"]), "cellr"), P(cm(x["remaining"]), "cellr")]
            if prior:
                row.append(P(_m(x["prior_ytd"]), "cellr"))
        else:
            row = [P(label, "cell"), P(cm(x["columns"][0]), "cellr"), P(cm(x["budgeted"]), "cellr"),
                   P(cm(x["ytd"]), "cellr"), P(cm(x["remaining"]), "cellr")]
            if prior:
                row += [P(_m(x["prior_period"]), "cellr"), P(_m(x["prior_ytd"]), "cellr")]
        return row

    for r in sec["rows"]:
        data.append(cells(r, r["label"] + (f" ({r['status']})" if r["raw_status"] in ("REJECTED", "INACTIVE") else "")))
        for c in r["children"]:
            data.append(cells(c, "    " + c["label"]))
    if not sec["rows"]:
        data.append([P(f"No {key} budgets.", "cell")] + [""] * (len(head) - 1))
    data.append(cells(sec["total"], "Total"))
    n = len(head) - 1
    col = 0.7 * inch if n > 6 else 0.86 * inch
    return [_grid(data, [W - n * col] + [col] * n), Paragraph(f"<i>{escape(GROSS_BUDGET_NOTE)}</i>", S["small"])]


def build_pdf(org: str, rep: dict) -> tuple[bytes, _AuditDoc]:
    buf = io.BytesIO()
    doc = _AuditDoc(buf, title=title(rep), author=org)
    W = FRAME_W - 12
    band = Table([[[P(org, "fs_w"), Paragraph(escape(title(rep)), S["fs_wt"]), P(as_of(rep), "fs_wb"),
                    P(f"Period {_d(rep['start_date'])} – {_d(rep['end_date'])}", "fs_w")]]], colWidths=[W])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _band()), ("LEFTPADDING", (0, 0), (-1, -1), 12),
                              ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    f: list = [_Mark(doc, rep["period_label"]), band, Spacer(1, 8), *heading("Highlights")]
    hl = highlight_tiles(rep)
    tt = Table([[[P(k, "fs_tk"), P(v, "fs_tv"), P(c, "small")] for k, v, c in hl]], colWidths=[W / len(hl)] * len(hl))
    tt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _tile()),
                            ("BOX", (0, 0), (-1, -1), 0.5, branding.color("outline_variant")),
                            ("INNERGRID", (0, 0), (-1, -1), 0.5, branding.color("outline_variant")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    f += [tt, Spacer(1, 6)]
    if rep["commentary"]:
        f += [*heading("Commentary", "h2")]
        f += [P(par, "body") for par in rep["commentary"].split("\n") if par.strip()]
    # accounts
    f += [Spacer(1, 8), *heading("Accounts")]
    acc = rep["accounts"]
    prior = rep["prior_fiscal_year"] is not None
    if rep["annual"]:
        head = ["Account", "Start of year", *[f"{q} end" for q in rep["quarter_names"]], "End"] + (["Prior year end"] if prior else [])
    else:
        head = ["Account", "Start", "End", "Change"] + ([f"Prior {rep['period']} end"] if prior else [])
    data = [_hdr(head)]
    for a in acc:
        name = f"{a['name']} ({a['number']})" + (" — closed" if a["closed"] else "")
        if rep["annual"]:
            row = [P(name, "cell"), P(_m(a["start"]), "cellr"), *[P(_m(v), "cellr") for v in a["quarter_ends"]],
                   P(_m(a["end"]), "cellr")]
        else:
            row = [P(name, "cell"), P(_m(a["start"]), "cellr"), P(_m(a["end"]), "cellr"), P(_m(a["change"]), "cellr")]
        if prior:
            row.append(P(_m(a["prior_end"]), "cellr"))
        data.append(row)
    tot_s = sum(a["start"] or 0 for a in acc)
    tot_e = sum(a["end"] or 0 for a in acc)
    if rep["annual"]:
        qsum = [sum((a["quarter_ends"][i] or 0) for a in acc) for i in range(len(rep["quarter_names"]))]
        total = [P("Total", "cellb"), P(cm(tot_s), "cellr"), *[P(cm(v), "cellr") for v in qsum], P(cm(tot_e), "cellr")]
    else:
        total = [P("Total", "cellb"), P(cm(tot_s), "cellr"), P(cm(tot_e), "cellr"), P(cm(tot_e - tot_s), "cellr")]
    if prior:
        total.append(P(cm(sum(a["prior_end"] or 0 for a in acc)), "cellr"))
    data.append(total)
    n = len(head) - 1
    f.append(_grid(data, [W - n * 0.95 * inch] + [0.95 * inch] * n))
    f.append(P("Account numbers are masked. Balances are by transaction date (cleared or not); a non-register account "
               "uses its balance history. — means no balance was recorded for that date.", "small"))
    # cash flow
    cf = rep["cash_flow"]
    f += [Spacer(1, 8), *heading("Cash flow")]
    data = [_hdr(["Month", "Money in", "Money out", "Net"])]
    for m in cf["months"]:
        data.append([P(m["label"], "cell"), P(cm(m["in"]), "cellr"), P(cm(m["out"]), "cellr"), P(cm(m["net"]), "cellr")])
    p_ = cf["period"]
    data.append([P("Total " + ("for the year" if rep["annual"] else rep["period"]), "cellb"), P(cm(p_["in"]), "cellr"),
                 P(cm(p_["out"]), "cellr"), P(cm(p_["net"]), "cellr")])
    f.append(_grid(data, [W - 3 * 1.3 * inch] + [1.3 * inch] * 3))
    rows = []
    if rep["annual"]:
        rows += [[P(q["label"], "cell"), P(cm(q["in"]), "cellr"), P(cm(q["out"]), "cellr"), P(cm(q["net"]), "cellr")]
                 for q in cf["quarters"]]
    else:
        y = cf["ytd"]
        rows.append([P("Year to date", "cell"), P(cm(y["in"]), "cellr"), P(cm(y["out"]), "cellr"), P(cm(y["net"]), "cellr")])
    if "prior_period" in cf:
        pp = cf["prior_period"]
        rows.append([P(f"Prior year: {pp['label']}", "cell"), P(cm(pp["in"]), "cellr"), P(cm(pp["out"]), "cellr"),
                     P(cm(pp["net"]), "cellr")])
        if not rep["annual"]:
            py = cf["prior_ytd"]
            rows.append([P("Prior year to date", "cell"), P(cm(py["in"]), "cellr"), P(cm(py["out"]), "cellr"),
                         P(cm(py["net"]), "cellr")])
    if rows:
        f += [Spacer(1, 4), _table([_hdr(["Summary", "Money in", "Money out", "Net"])] + rows,
                                   [W - 3 * 1.3 * inch] + [1.3 * inch] * 3)]
    # budgets
    for key, label in (("income", "Income"), ("expense", "Expense")):
        f += [Spacer(1, 8), *heading(f"{label} budgets", "h2")]
        if key == "income":
            f.append(P("The Fiscal Year's budgets with their sub-budgets (Rejected or Inactive ones only when they have "
                       "activity). Over / under = budgeted − actual; negative = over budget.", "small"))
        f += _budget_table(rep, key, W)
    if rep["budgets"]["budget_zero"]:
        f.append(P(f"Budget 0 (transfers and other non-budget activity) year to date: "
                   f"{cm(rep['budgets']['budget_zero']['ytd'])}; not counted above.", "small"))
    # fundraisers
    fr = rep["fundraisers"]
    if fr["enabled"]:
        f += [Spacer(1, 8), *heading("Fundraisers")]
        listed = fr["listed"]
        if listed:
            data = [_hdr(["Fundraiser", "Ended", "Quarter", "Gross", "Net", "Note"])]
            for r in listed:
                note = []
                if r["revised"]:
                    note.append(f"Revised: {cm(r['change'])} net in {rep['period']}")
                if r["after_cutoff"]:
                    note.append(f"Activity after {_d(rep['cutoff'])} not included")
                data.append([P(r["name"], "cell"), P(_d(r["end_date"]), "cell"), P(r["quarter"] or "", "cell"),
                             P(cm(r["gross"]), "cellr"), P(cm(r["net"]), "cellr"), P("; ".join(note), "small")])
            f.append(_table(data, [1.8 * inch, 1.15 * inch, 0.6 * inch, 0.95 * inch, 0.95 * inch, W - 5.45 * inch]))
        else:
            f.append(P("No fundraiser ended in this period.", "body"))
        if not rep["annual"] and fr["rows"]:
            ytd = [r for r in fr["rows"]]
            f.append(P(f"Fiscal Year to date: {len(ytd)} fundraiser{'s' if len(ytd) != 1 else ''}, gross "
                       f"{cm(fr['total_gross'])}, net {cm(fr['total_net'])}.", "small"))
        if rep["annual"]:
            share = "" if fr["share_of_income"] is None else f", {fr['share_of_income']:.1f}% of the year's income"
            f.append(P(f"Fundraising this year: gross {cm(fr['total_gross'])}, net {cm(fr['total_net'])}{share}.", "body"))
        for d in fr["details"]:
            f += [*heading(d["name"], "h2")]
            data = [_hdr(["Income by budget line", "Amount"])] + [[P(k, "cell"), P(cm(v), "cellr")] for k, v in d["income_lines"]]
            data.append([P("Gross", "cellb"), P(cm(d["gross"]), "cellr")])
            f.append(_grid(data, [W - 1.3 * inch, 1.3 * inch]))
            data = [_hdr(["Expenses by budget line", "Amount"])] + [[P(k, "cell"), P(cm(v), "cellr")] for k, v in d["expense_lines"]]
            data.append([P("Expenses", "cellb"), P(cm(d["expense"]), "cellr")])
            f += [Spacer(1, 3), _grid(data, [W - 1.3 * inch, 1.3 * inch])]
            lines = [f"Net {cm(d['net'])}."]
            if d["seed_out"] or d["seed_back"]:
                lines.append(f"Seed cash: {cm(d['seed_out'])} out, {cm(d['seed_back'])} returned, net {cm(d['seed_net'])} "
                             "(not counted in gross or expenses).")
            py = d["prior_year"]
            if py:
                lines.append(f"Last year ({py['label']}): gross {_m(py['gross'])}, net {cm(py['net'])}"
                             f"{' (estimated)' if py.get('estimated') else ''}; change in net {_pct(py['net_change_pct'])}"
                             + ("" if py["gross"] is None else f", in gross {_pct(py['gross_change_pct'])}") + ".")
            g = d["goal"]
            if g:
                lines.append(f"Goal {cm(int(round(float(g['goal']) * 100)))}: net {'met' if g['met'] else 'below'} the goal"
                             + ("" if g["pct_of_goal"] is None else f" ({g['pct_of_goal']:.1f}% of goal)") + ".")
            if d["after_cutoff"]:
                lines.append(f"Activity after {_d(rep['cutoff'])} not included.")
            f += [P(x, "body") for x in lines]
    f += [Spacer(1, 10)] + [P(x, "small") for x in notes_lines(rep)]
    doc.build(f)
    return buf.getvalue(), doc


def pdf_file(org: str, rep: dict) -> tuple[str, str, int]:
    pdf, doc = build_pdf(org, rep)
    footer = f"{org}  ·  {title(rep)}  ·  {as_of(rep)}"
    fd, path = tempfile.mkstemp(prefix="fmpoc-period-", suffix=".pdf")
    os.close(fd)
    try:
        pages = _stamp_and_write(pdf, doc, path, footer)
    except Exception:
        os.unlink(path)
        raise
    return path, file_stem(org, rep) + "-report.pdf", pages


# ------------------------------------------------------------------ presentation
def build_pptx(org: str, rep: dict) -> bytes:
    prs = Presentation()
    deck._strip_template_extras(prs)
    prs.slide_width, prs.slide_height = deck.W, deck.H
    blank = prs.slide_layouts[6]
    # 1. title + highlights (commentary in the notes)
    sl = prs.slides.add_slide(blank)
    deck._mark(sl, Inches(0.6), Inches(0.9), Inches(1.0), Inches(1.3))
    deck._text(sl, Inches(2.0), Inches(0.8), Inches(10.5), Inches(0.6), org, size=24, bold=True)
    deck._text(sl, Inches(2.0), Inches(1.35), Inches(10.5), Inches(0.9), title(rep), size=34, bold=True, color=deck.C.BLUE)
    deck._text(sl, Inches(2.0), Inches(2.25), Inches(10.5), Inches(0.5), as_of(rep), size=18,
               color=(deck.C.GOLD if deck.C.DARK else deck.RED) if rep["preliminary"] else deck.C.GREY)
    hl = highlight_tiles(rep)
    w = 12.1 / len(hl)
    for i, (k, v, c) in enumerate(hl):
        x = Inches(0.6 + i * w)
        card = sl.shapes.add_shape(1, x, Inches(3.4), Inches(w - 0.25), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = deck.C.LIGHT
        card.line.color.rgb = deck.C.LINE
        deck._text(sl, x + Inches(0.2), Inches(3.55), Inches(w - 0.6), Inches(0.5), k, size=15, bold=True, color=deck.C.GREY)
        deck._text(sl, x + Inches(0.2), Inches(4.1), Inches(w - 0.6), Inches(0.9), v, size=24, bold=True)
        deck._text(sl, x + Inches(0.2), Inches(5.05), Inches(w - 0.6), Inches(0.6), c, size=12, color=deck.C.GREY)
    deck._notes(sl, ([f"Commentary: {rep['commentary']}"] if rep["commentary"] else []) + notes_lines(rep))
    # 2. accounts
    acc = rep["accounts"]
    rows = [[f"{a['name']} ({a['number']})", _m(a["start"]), _m(a["end"]), _m(a["change"]),
             *([_m(a["prior_end"])] if rep["prior_fiscal_year"] else [])] for a in acc]
    tot_s, tot_e = sum(a["start"] or 0 for a in acc), sum(a["end"] or 0 for a in acc)
    rows.append(["Total", cm(tot_s), cm(tot_e), cm(tot_e - tot_s),
                 *([cm(sum(a["prior_end"] or 0 for a in acc))] if rep["prior_fiscal_year"] else [])])
    head = ["Account", "Start", "End", "Change"] + (["Prior year"] if rep["prior_fiscal_year"] else [])
    for k, part in enumerate(deck._chunks(rows, deck.TABLE_ROWS)):
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Accounts" + (" (continued)" if k else ""), f"Start and end of {'the year' if rep['annual'] else rep['period']}")
        widths = [4.9, 1.8, 1.8, 1.8] + ([1.8] if rep["prior_fiscal_year"] else [])
        deck._table(sl, Inches(0.6), Inches(1.7), Inches(sum(widths)), [head] + part, widths, set(range(1, len(head))),
                    bold_rows=[i + 1 for i, r in enumerate(part) if r[0] == "Total"])
        deck._notes(sl, ["Account numbers are masked. Balances by transaction date as of the cutoff."])
    # 3. cash flow chart: monthly money in / out bars and the net line
    cf = rep["cash_flow"]
    sl = prs.slides.add_slide(blank)
    deck._title(sl, "Cash flow", "Money in and out by month" + ("" if rep["annual"] else f" · YTD net {cm(cf['ytd']['net'])}"))
    cd = CategoryChartData(number_format='"$"#,##0')
    cd.categories = [m["label"] for m in cf["months"]] or ["—"]
    cd.add_series("Money in", [m["in"] / 100 for m in cf["months"]] or [0])
    cd.add_series("Money out", [m["out"] / 100 for m in cf["months"]] or [0])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.0), cd)
    ch = gf.chart
    deck._style_chart(ch, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM)
    for i, s in enumerate(ch.plots[0].series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = deck.C.SERIES[i % len(deck.C.SERIES)]
    _add_net_line(ch, [m["net"] / 100 for m in cf["months"]] or [0])
    lines = [f"{m['label']}: in {cm(m['in'])}, out {cm(m['out'])}, net {cm(m['net'])}." for m in cf["months"]]
    lines.append(f"{'Year' if rep['annual'] else rep['period']} total: net {cm(cf['period']['net'])}; year to date net {cm(cf['ytd']['net'])}.")
    if rep["annual"]:
        lines += [f"{q['label']}: net {cm(q['net'])}." for q in cf["quarters"]]
    if "prior_period" in cf:
        lines.append(f"Prior year ({cf['prior_period']['label']}): net {cm(cf['prior_period']['net'])}.")
    deck._notes(sl, lines + ["Transfers between accounts (Budget 0) are left out."])
    # 4. budgets (tables)
    b = rep["budgets"]
    for key, label in (("income", "Income"), ("expense", "Expense")):
        sec = b[key]
        rows = []
        for r in sec["rows"]:
            for x, ind in [(r, ""), *[(c, "   ") for c in r["children"]]]:
                rows.append([ind + x["label"], *([cm(v) for v in x["columns"]]), cm(x["ytd"]), cm(x["budgeted"]), cm(x["remaining"])])
        t = sec["total"]
        rows.append(["Total", *[cm(v) for v in t["columns"]], cm(t["ytd"]), cm(t["budgeted"]), cm(t["remaining"])])
        head = ["Budget", *(b["columns"] if rep["annual"] else ["Quarter"]), "YTD actual" if not rep["annual"] else "Actual",
                "Budgeted", "Over / under"]
        n = len(head) - 1
        widths = [12.1 - n * 1.45] + [1.45] * n
        for k, part in enumerate(deck._chunks(rows, deck.TABLE_ROWS)):
            sl = prs.slides.add_slide(blank)
            deck._title(sl, f"{label} budgets" + (" (continued)" if k else ""), f"Year to date as of {_d(rep['cutoff'])}")
            deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [head] + part, widths, set(range(1, len(head))),
                        bold_rows=[i + 1 for i, r in enumerate(part) if r[0] == "Total"], size=11)
            deck._text(sl, Inches(0.6), Inches(6.75), Inches(12.1), Inches(0.5), GROSS_BUDGET_NOTE, size=11, color=deck.C.GREY)
            deck._notes(sl, ["Over / under = budgeted − actual; negative = over budget.", GROSS_BUDGET_NOTE])
    # 5. fundraisers
    fr = rep["fundraisers"]
    if fr["enabled"]:
        rows = [[r["name"] + (" (Revised)" if r["revised"] else ""), r["quarter"] or "", cm(r["gross"]), cm(r["net"])]
                for r in fr["listed"]]
        rows.append(["Total" if rep["annual"] else f"Ended in {rep['period']}",
                     "", cm(sum(r["gross"] for r in fr["listed"] if r["in_period"])),
                     cm(sum(r["net"] for r in fr["listed"] if r["in_period"]))])
        sl = prs.slides.add_slide(blank)
        deck._title(sl, "Fundraisers", f"Year to date: gross {cm(fr['total_gross'])}, net {cm(fr['total_net'])}"
                    + ("" if fr.get("share_of_income") is None or not rep["annual"] else f" · {fr['share_of_income']:.1f}% of income"))
        deck._table(sl, Inches(0.6), Inches(1.7), Inches(12.1), [["Fundraiser", "Quarter", "Gross", "Net"]] + rows[:deck.TABLE_ROWS],
                    [6.7, 1.4, 2.0, 2.0], {2, 3}, bold_rows=[len(rows[:deck.TABLE_ROWS])])
        deck._notes(sl, ["A fundraiser belongs to the quarter its event ends in. Revised = an earlier quarter's fundraiser "
                         f"with transactions dated in {rep['period']}." if not rep["annual"] else
                         "Gross = income counted for the fundraiser; net = gross − expenses."]
                    + [f"{r['name']}: activity after {_d(rep['cutoff'])} not included." for r in fr["listed"] if r["after_cutoff"]])
        for d in fr["details"]:
            sl = prs.slides.add_slide(blank)
            deck._title(sl, d["name"], f"Ended {_d(d['end_date'])} · gross {cm(d['gross'])} · net {cm(d['net'])}")
            rows = [[k, cm(v)] for k, v in d["income_lines"]] + [["Gross", cm(d["gross"])]]
            deck._table(sl, Inches(0.6), Inches(1.7), Inches(5.9), [["Income", "Amount"]] + rows[:10], [4.1, 1.8], {1},
                        bold_rows=[len(rows[:10])], size=11)
            rows = [[k, cm(v)] for k, v in d["expense_lines"]] + [["Expenses", cm(d["expense"])]]
            deck._table(sl, Inches(6.8), Inches(1.7), Inches(5.9), [["Expenses", "Amount"]] + rows[:10], [4.1, 1.8], {1},
                        bold_rows=[len(rows[:10])], size=11)
            extra = []
            if d["seed_out"] or d["seed_back"]:
                extra.append(f"Seed cash: {cm(d['seed_out'])} out, {cm(d['seed_back'])} returned, net {cm(d['seed_net'])}.")
            if d["prior_year"]:
                py = d["prior_year"]
                extra.append(f"Last year ({py['label']}): net {cm(py['net'])}; change {_pct(py['net_change_pct'])}.")
            if d["goal"]:
                g = d["goal"]
                extra.append(f"Goal {cm(int(round(float(g['goal']) * 100)))}: {'met' if g['met'] else 'not met'}"
                             + ("" if g["pct_of_goal"] is None else f" ({g['pct_of_goal']:.1f}%)") + ".")
            if extra:
                deck._text(sl, Inches(0.6), Inches(6.2), Inches(12.1), Inches(1.0), extra, size=14, color=deck.C.TEXT)
            deck._notes(sl, extra or [f"Net {cm(d['net'])}."])
    # closing
    sl = prs.slides.add_slide(blank)
    deck._text(sl, Inches(0.6), Inches(3.0), Inches(12.1), Inches(1.2), "Questions?", size=48, bold=True, align=PP_ALIGN.CENTER)
    deck._text(sl, Inches(0.6), Inches(4.3), Inches(12.1), Inches(0.6), f"{org} · {title(rep)} · {as_of(rep)}", size=18,
               color=deck.C.GREY, align=PP_ALIGN.CENTER)
    deck._notes(sl, [f"Generated {_d(rep['generated_on'])}."])
    prs.core_properties.title = title(rep)
    prs.core_properties.author = org
    deck.brand_footers(prs)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _add_net_line(chart, values: list[float]) -> None:
    """The net as a line over the monthly bars (a second plot in the same chart; python-pptx has no API for a combo
    chart, so the line chart XML is added beside the bar chart and shares its axes)."""
    from copy import deepcopy
    from lxml import etree
    from pptx.oxml.ns import qn
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    bar = plot_area.find(qn("c:barChart"))
    ser = bar.findall(qn("c:ser"))[0]
    line = etree.SubElement(plot_area, qn("c:lineChart"))
    plot_area.remove(line)
    bar.addnext(line)
    etree.SubElement(line, qn("c:grouping")).set("val", "standard")
    etree.SubElement(line, qn("c:varyColors")).set("val", "0")
    s = deepcopy(ser)
    for tag in ("c:invertIfNegative", "c:dLbls", "c:spPr"):
        for el in s.findall(qn(tag)):
            s.remove(el)
    s.find(qn("c:idx")).set("val", "2")
    s.find(qn("c:order")).set("val", "2")
    tx = s.find(qn("c:tx"))
    if tx is not None:
        for v in tx.iter(qn("c:v")):
            v.text = "Net"
    val = s.find(qn("c:val"))
    for i, pt in enumerate(val.iter(qn("c:pt"))):
        if i < len(values):
            pt.find(qn("c:v")).text = repr(float(values[i]))
    etree.SubElement(s, qn("c:smooth")).set("val", "0")
    marker = etree.Element(qn("c:marker"))
    etree.SubElement(marker, qn("c:symbol")).set("val", "circle")
    s.insert(list(s).index(s.find(qn("c:cat"))) if s.find(qn("c:cat")) is not None else len(s), marker)
    line.append(s)
    etree.SubElement(line, qn("c:marker")).set("val", "1")
    for ax in bar.findall(qn("c:axId")):
        etree.SubElement(line, qn("c:axId")).set("val", ax.get("val"))
