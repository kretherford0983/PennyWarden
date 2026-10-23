"""2.1.0 (#292): a series can be started on a single fundraiser, so results of years before PennyWarden can be entered
for an event with nothing in the app to "continue from" (e.g. the organization's first year in PennyWarden), and the
year over year of its reports and presentation picks them up."""
from __future__ import annotations

import io

from pptx import Presentation

from test_v161_fundraiser_manage import gala  # noqa: F401  (fixture)
from test_v210_fundraiser_series import pdf_text, years  # noqa: F401  (fixture)


def deck_titles(content: bytes) -> list[str]:
    return [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text), "")
            for sl in Presentation(io.BytesIO(content)).slides]


def test_first_year_in_pennywarden_prior_years_reach_the_reports(env, gala):
    gid = gala["id"]
    assert env.bm.get("/api/fundraiser-series").json() == []                      # nothing to continue from
    assert [o["id"] for o in env.bm.get("/api/fundraiser-series/start-options").json()] == [gid]
    r = env.bm.post(f"/api/fundraisers/{gid}/series")
    assert r.status_code == 200, r.text
    sid = r.json()["series_id"]
    s = r.json()["series"]
    assert [(x["id"], x["name"], [f["id"] for f in x["fundraisers"]]) for x in s] == [(sid, "Gala", [gid])]
    assert env.bm.get(f"/api/fundraisers/{gid}").json()["series_link"]["series"] == {"id": sid, "name": "Gala"}
    assert env.bm.get("/api/fundraiser-series/start-options").json() == []
    # a series with no history adds nothing to the report yet
    assert "Year over year" not in pdf_text(env.bu.get(f"/api/fundraisers/{gid}/summary-report").content)
    for label, net in (("FY2024", "700.00"), ("FY2025", "900.00")):
        assert env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": label, "net": net}).status_code == 201
    t = pdf_text(env.bu.get(f"/api/fundraisers/{gid}/summary-report").content)
    assert "Year over year" in t and "FY2024 (Entered)" in t and "FY2025 (Entered)" in t
    assert "Year over year" in deck_titles(env.bu.get(f"/api/fundraisers/{gid}/presentation").content)


def test_start_is_idempotent_and_audited(env, years):
    rid = years["raffle"]["id"]
    assert [o["name"] for o in env.bm.get("/api/fundraiser-series/start-options").json()] == ["Raffle"]   # golf is linked
    a = env.bm.post(f"/api/fundraisers/{rid}/series").json()["series_id"]
    b = env.bm.post(f"/api/fundraisers/{rid}/series").json()["series_id"]
    assert a == b
    ev = env.auditor.get("/api/audit-events?action=FUNDRAISER_SERIES_CREATED").json()["items"]
    assert sum(1 for e in ev if (e.get("after") or {}).get("fundraiser_ids") == [rid]) == 1
    link = env.auditor.get("/api/audit-events?action=FUNDRAISER_SERIES_LINK_UPDATED").json()["items"]
    assert any((e.get("after") or {}).get("series_id") == a for e in link)
    # a fundraiser already in a series gets that series, not a new one
    golf = env.bm.get(f"/api/fundraisers/{years['g27']['id']}").json()["series_link"]["series"]["id"]
    assert env.bm.post(f"/api/fundraisers/{years['g27']['id']}/series").json()["series_id"] == golf
    assert sorted(x["name"] for x in env.bm.get("/api/fundraiser-series").json()) == ["Golf Outing", "Raffle"]


def test_year_in_the_app_still_refused(env, years):
    sid = env.bm.post(f"/api/fundraisers/{years['raffle']['id']}/series").json()["series_id"]
    r = env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2027", "net": "1.00"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "YEAR_RECORDED"
    assert env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2026", "net": "1.00"}).status_code == 201


def test_clearing_continues_from_keeps_a_started_series(env, years):
    rid = years["raffle"]["id"]
    sid = env.bm.post(f"/api/fundraisers/{rid}/series").json()["series_id"]
    r = env.bm.put(f"/api/fundraisers/{rid}/series-link", {"continues_from_id": None})
    assert r.status_code == 200 and r.json()["series_link"]["series"]["id"] == sid
    # an ordinary edit leaves it too
    f = years["raffle"]
    body = {"name": f["name"], "description": "Spring raffle", "start_date": f["start_date"], "end_date": None,
            "budget_ids": [b["id"] for b in f["budgets"]], "filter_text": "Raffle", "filter_regex": False,
            "continues_from_id": None}                                       # as the edit form sends it
    r = env.bm.put(f"/api/fundraisers/{rid}", body)
    assert r.status_code == 200, r.text
    assert env.bm.get(f"/api/fundraisers/{rid}").json()["series_link"]["series"]["id"] == sid
    # a fundraiser that continues from another still leaves its series when the link is removed
    r = env.bm.put(f"/api/fundraisers/{years['g27']['id']}/series-link", {"continues_from_id": None})
    assert r.status_code == 200 and r.json()["series_link"]["series"] is None


def test_access(env, years):
    rid = years["raffle"]["id"]
    for c in (env.bu, env.ru, env.auditor, env.admin):
        assert c.post(f"/api/fundraisers/{rid}/series").status_code == 403
        assert c.get("/api/fundraiser-series/start-options").status_code == 403
    assert env.admin.put("/api/system/modules", {"fundraisers": False}).status_code == 200
    assert env.bm.post(f"/api/fundraisers/{rid}/series").status_code == 404
