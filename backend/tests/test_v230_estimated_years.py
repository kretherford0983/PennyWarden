"""2.3.0 build 5 (#390): entered years before PennyWarden can be marked estimated - fundraiser series results and budget
years. Decisions (Kyle, 2026-10-29): an Estimated box; the note becomes required when it is ticked (server-validated);
a Correct action on entered fundraiser results (reason required, audited before/after, the flag included); reports
and presentations mark the year with * and a footnote "* Estimated: <note>"; charts mark the label (and use a lighter
bar), never color alone (BR-079); the budget History, Before PennyWarden, the CSV import (optional estimated column)
and the Budget Review (screen, PDF, CSV) carry the same mark."""
from __future__ import annotations

import csv
import io

from pypdf import PdfReader

AS_OF = "2026-09-30"


def _text(content: bytes) -> str:
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(io.BytesIO(content)).pages)


def _enable(env):
    assert env.admin.put("/api/system/modules", {"fundraisers": True}).status_code == 200


def _gala_series(env, base):
    """A fundraiser in FY2027 that starts a series; returns (fundraiser, series id)."""
    _enable(env)
    f = env.bm.post("/api/fundraisers", {"name": "Gala", "start_date": "2026-09-12",
                                         "budget_ids": [base["inc"]["id"]]}).json()
    r = env.bm.post(f"/api/fundraisers/{f['id']}/series", {})
    assert r.status_code in (200, 201), r.text
    return f, r.json()["series_id"]


def _result(env, sid, label, net, **kw):
    return env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": label, "net": net, **kw})


def _results(env, sid):
    s = next(x for x in env.bm.get("/api/fundraiser-series").json() if x["id"] == sid)
    return s["results"]


# ------------------------------------------------------------------ fundraiser series
def test_estimated_result_needs_a_note_and_is_marked(env, base):
    f, sid = _gala_series(env, base)
    r = _result(env, sid, "FY2024", "1200.00", estimated=True)
    assert r.status_code == 422 and r.json()["error"]["errors"][0]["field"] == "note"
    assert _result(env, sid, "FY2024", "1200.00", estimated=True, note="records lost; from memory").status_code == 201
    assert _result(env, sid, "FY2025", "1500.00").status_code == 201                   # not estimated: note optional
    rs = {r["fiscal_year_label"]: r for r in _results(env, sid)}
    assert rs["FY2024"]["estimated"] is True and rs["FY2025"]["estimated"] is False
    # the stakeholder report and its PDF: FY2024* with the footnote
    pdf = env.bm.get(f"/api/fundraisers/{f['id']}/summary-report")
    assert pdf.status_code == 200, pdf.text
    text = _text(pdf.content)
    assert "2024*" in text and "* Estimated: 2024 — records lost; from memory" in text


def test_correct_entered_result_reason_required_and_audited(env, base):
    f, sid = _gala_series(env, base)
    _result(env, sid, "FY2024", "1200.00")
    rid = _results(env, sid)[0]["id"]
    url = f"/api/fundraiser-series/{sid}/results/{rid}/correct"
    assert env.bm.post(url, {"net": "1250.00"}).status_code == 422                       # reason required
    assert env.bm.post(url, {"net": "1250.00", "estimated": True, "reason": "x"}).status_code == 422   # note required
    r = env.bm.post(url, {"net": "1250.00", "estimated": True, "note": "partly from bank statements",
                          "reason": "found the deposit slips"})
    assert r.status_code == 200, r.text
    x = _results(env, sid)[0]
    assert (x["net"], x["estimated"], x["note"]) == ("1250.00", True, "partly from bank statements")
    ev = env.auditor.get("/api/audit-events?action=FUNDRAISER_RESULT_CORRECTED").json()["items"][0]
    assert ev["before"]["estimated"] is False and ev["after"]["estimated"] is True
    assert ev["before"]["net"] == "1200.00" and ev["after"]["reason"] == "found the deposit slips"
    # un-marking it
    r = env.bm.post(url, {"net": "1250.00", "estimated": False, "note": "", "reason": "confirmed"})
    assert r.status_code == 200 and _results(env, sid)[0]["estimated"] is False
    # permissions: only Budget Managers; retired results can't be corrected
    for c in (env.bu, env.ru, env.auditor, env.admin):
        assert c.post(url, {"net": "1.00", "reason": "x"}).status_code == 403
    assert env.bm.post(f"/api/fundraiser-series/{sid}/results/{rid}/retire", {"reason": "dup"}).status_code == 200
    assert env.bm.post(url, {"net": "1.00", "reason": "x"}).status_code == 409


def test_fiscal_year_fundraising_report_marks_estimated_years(env, base):
    f, sid = _gala_series(env, base)
    _result(env, sid, "FY2025", "900.00", estimated=True, note="bundled with the raffle")
    rep = env.bm.get("/api/fundraisers/fy-report", params={"fiscal_year_id": base["fy"]["id"], "ids": f["id"], "format": "json"})
    assert rep.status_code == 200, rep.text
    d = rep.json()
    h = next(x for x in d["history"] if x["year"] == 2025)
    assert h["estimated"] is True
    assert next(t for t in d["trend"] if t["year"] == 2025)["estimated"] is True
    assert "* Estimated: 2025 Gala — bundled with the raffle" in d["footnotes"]
    pdf = env.bm.get("/api/fundraisers/fy-report", params={"fiscal_year_id": base["fy"]["id"], "ids": f["id"], "format": "pdf"})
    assert pdf.status_code == 200
    assert "2025*" in _text(pdf.content)
    # the presentations build with an estimated year (lighter bar, * label, footnote in the notes)
    from pptx import Presentation
    deck = env.bm.get("/api/fundraisers/fy-report", params={"fiscal_year_id": base["fy"]["id"], "ids": f["id"], "format": "pptx"})
    assert deck.status_code == 200
    notes = " ".join(sl.notes_slide.notes_text_frame.text for sl in Presentation(io.BytesIO(deck.content)).slides
                     if sl.has_notes_slide)
    assert "* Estimated: 2025 Gala — bundled with the raffle" in notes
    one = env.bm.get(f"/api/fundraisers/{f['id']}/presentation")
    assert one.status_code == 200, one.text
    notes = " ".join(sl.notes_slide.notes_text_frame.text for sl in Presentation(io.BytesIO(one.content)).slides
                     if sl.has_notes_slide)
    assert "* Estimated: 2025 — bundled with the raffle" in notes
    docx = env.bm.get(f"/api/fundraisers/{f['id']}/summary-report", params={"format": "docx"})
    assert docx.status_code == 200


# ------------------------------------------------------------------ budget years
def _former(env, code="450", name="Bowling"):
    return env.bm.post("/api/budgets/former", {"code": code, "name": name, "budget_type": "EXPENSE"}).json()


def test_budget_year_estimated_note_required_and_review_marks_it(env, base):
    f = _former(env)
    url = f"/api/budgets/former/{f['id']}/entries"
    r = env.bm.post(url, {"fiscal_year_label": "FY2024", "budgeted": "800", "actual": "760", "estimated": True})
    assert r.status_code == 422 and r.json()["error"]["errors"][0]["field"] == "note"
    r = env.bm.post(url, {"fiscal_year_label": "FY2024", "budgeted": "800", "actual": "760", "estimated": True,
                          "note": "the old treasurer's guess"})
    assert r.status_code == 201, r.text
    e = r.json()
    assert e["estimated"] is True
    # Correct can clear it; clearing the note while estimated is refused
    assert env.bm.patch(f"/api/budgets/history-entries/{e['id']}", {"note": "", "reason": "x"}).status_code == 422
    rep = env.bm.get("/api/reports/budget-review", params={"fiscal_year_id": base["fy"]["id"], "as_of": AS_OF, "years": 5}).json()
    row = next(x for x in rep["expense"] if x["code"] == "450")
    cell = row["history"][0]
    assert cell["estimated"] is True and cell["estimate_note"] == "the old treasurer's guess"
    c = env.bm.get("/api/reports/budget-review", params={"fiscal_year_id": base["fy"]["id"], "as_of": AS_OF, "years": 5,
                                                        "format": "csv"})
    line = next(x for x in csv.reader(io.StringIO(c.text)) if len(x) > 2 and x[2] == "450")
    assert any("estimated: the old treasurer's guess" in v for v in line)
    p = env.bm.get("/api/reports/budget-review", params={"fiscal_year_id": base["fy"]["id"], "as_of": AS_OF, "years": 5,
                                                        "format": "pdf"})
    t = _text(p.content)
    assert "* Estimated: FY2024 450 Bowling — the old treasurer's guess" in t
    ev = env.auditor.get("/api/audit-events?action=BUDGET_HISTORY_ENTERED").json()["items"][0]
    assert ev["after"]["estimated"] is True


def test_csv_import_estimated_column(env, base):
    text = ("budget code,type,former code,former name,year,budgeted,actual,note,estimated\n"
            ",Expense,450,Bowling,FY2023,800,760,,yes\n")
    p = env.bm.post("/api/budgets/history-entries/preview", {"fiscal_year_id": base["fy"]["id"], "csv": text}).json()
    assert p["counts"]["REFUSED"] == 1                                       # estimated without a note
    text = text.replace(",,yes", ",guessed from minutes,yes") + ",Expense,450,Bowling,FY2024,900,910,,no\n"
    r = env.bm.post("/api/budgets/history-entries/import", {"fiscal_year_id": base["fy"]["id"], "csv": text})
    assert r.status_code == 200 and r.json()["saved"] == 2, r.text
    f = next(x for x in env.bm.get("/api/budgets/former").json() if x["label"].startswith("450"))
    d = env.bm.get(f"/api/budgets/former/{f['id']}").json()
    assert [(x["fiscal_year_label"], x["estimated"]) for x in d["entries"]] == [("FY2023", True), ("FY2024", False)]
    bad = text.replace(",no\n", ",maybe\n")
    p = env.bm.post("/api/budgets/history-entries/preview", {"fiscal_year_id": base["fy"]["id"], "csv": bad}).json()
    assert p["counts"]["REFUSED"] >= 1
