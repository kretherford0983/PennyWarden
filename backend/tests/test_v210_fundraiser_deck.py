"""2.1.0 (#276): the fundraiser presentation (.pptx) - native charts with data labels, speaker notes on every slide,
figures equal to the stakeholder report, no Entity names when names are off, and no macros, media, embedded fonts
or external links (PowerPoint, Google Slides and Keynote compatibility)."""
from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from test_v161_fundraiser_manage import gala, put_line  # noqa: F401  (fixture)
from test_v210_fundraiser_summary import buckets  # noqa: F401  (fixture)


def deck(env, g, client=None, **q):
    params = {k: str(v).lower() for k, v in q.items()}
    return (client or env.bu).get(f"/api/fundraisers/{g['id']}/presentation", params=params)


def slides(content: bytes):
    return Presentation(io.BytesIO(content)).slides


def all_text(content: bytes) -> str:
    prs = Presentation(io.BytesIO(content))
    out = []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if sh.has_table:
                out += [c.text for r in sh.table.rows for c in r.cells]
            if sh.has_chart:
                ch = sh.chart
                out += [str(c) for c in ch.plots[0].categories]
                out += [s.name for p in ch.plots for s in p.series]
        out.append(sl.notes_slide.notes_text_frame.text)
    return " ".join(out)


def charts(content: bytes) -> dict:
    out = {}
    for i, sl in enumerate(slides(content)):
        for sh in sl.shapes:
            if sh.has_chart:
                out.setdefault(sh.chart.chart_type, []).append((i, sh.chart))
    return out


def titles(content: bytes) -> list[str]:
    out = []
    for sl in slides(content):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text]
        out.append(texts[0] if texts else "")
    return out


def test_deck_slides_charts_and_notes(env, buckets):
    r = deck(env, buckets)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert r.headers["content-disposition"] == 'attachment; filename="fundraiser-Gala-presentation.pptx"'
    tl = titles(r.content)
    assert tl[0] == "Acme Org" or "Gala" in tl[0] or tl[0]
    for want in ("Executive highlights", "From revenue to proceeds", "Revenue by category", "Category detail",
                 "Net by category / bucket", "Revenue over time", "Expenses", "Key takeaways", "Questions?"):
        assert any(x.startswith(want) for x in tl), (want, tl)
    assert any(x.startswith("Top contributors") for x in tl)
    c = charts(r.content)
    assert XL_CHART_TYPE.DOUGHNUT in c and XL_CHART_TYPE.COLUMN_CLUSTERED in c
    assert XL_CHART_TYPE.LINE_MARKERS in c and XL_CHART_TYPE.BAR_CLUSTERED in c
    for kind in c.values():
        for _i, ch in kind:
            assert ch.plots[0].has_data_labels
    # chart data equals the report figures
    _i, col = c[XL_CHART_TYPE.COLUMN_CLUSTERED][0]
    assert list(col.plots[0].series[0].values) == [1325.0, 300.0, 1025.0]
    _i, dn = c[XL_CHART_TYPE.DOUGHNUT][0]
    assert sorted(dn.plots[0].series[0].values) == [75.0, 1250.0]
    # speaker notes on every slide, with the calculations
    prs = Presentation(io.BytesIO(r.content))
    assert all(sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip() for sl in prs.slides)
    text = all_text(r.content)
    assert "Net margin = net proceeds ÷ net revenue = 77.4%" in text
    assert "Less: Pre-paid team credit" in text and "$1,025.00" in text
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000          # 16:9


def test_compatibility_rules(env, buckets):
    z = zipfile.ZipFile(io.BytesIO(deck(env, buckets).content))
    names = z.namelist()
    assert not any(n.startswith("ppt/media/") for n in names)
    assert not any("vbaProject" in n or n.endswith(".bin") or "thumbnail" in n for n in names)   # no macros, no binaries
    assert all(n.endswith(".xlsx") for n in names if n.startswith("ppt/embeddings/"))       # only the chart data
    assert not any(n.startswith("ppt/fonts/") for n in names)
    rels = " ".join(z.read(n).decode("utf-8", "replace") for n in names if n.endswith(".rels"))
    assert 'TargetMode="External"' not in rels
    slides_xml = " ".join(z.read(n).decode() for n in names if n.startswith("ppt/slides/slide"))
    assert "<p:transition" not in slides_xml and "<p:timing" not in slides_xml
    assert 'typeface="Arial"' in slides_xml


def test_names_off_leaves_no_entity_names(env, buckets):
    ent = env.entity("Jordan Benefactor")
    env.txn(buckets["base"]["acct"]["id"], "DEPOSIT", [{"budget_id": buckets["base"]["inc_leaf"], "amount": "40.00",
                                                       "description": "Sponsor"}], date="2026-09-15", entity_id=ent["id"])
    line = next(x for x in env.bu.get(f"/api/fundraisers/{buckets['id']}").json()["lines"] if x["description"] == "Sponsor")
    env.ru.put(f"/api/fundraisers/{buckets['id']}/lines/{line['allocation_id']}",
               {"buckets": [{"bucket_id": buckets["ids"]["Raffle"], "amount": "40.00"}]})
    assert "Jordan Benefactor" in all_text(deck(env, buckets).content)
    r = deck(env, buckets, names=False)
    blob = b"".join(zipfile.ZipFile(io.BytesIO(r.content)).read(n) for n in zipfile.ZipFile(io.BytesIO(r.content)).namelist())
    assert b"Jordan Benefactor" not in blob                                   # nowhere: slides, notes, chart data
    assert not any(t.startswith("Top contributors") for t in titles(r.content))


def test_no_buckets_leaves_out_category_slides(env, gala):
    tl = titles(deck(env, gala).content)
    assert not any(x.startswith(("Revenue by category", "Category detail", "Net by category", "Top contributors")) for x in tl)
    assert any(x.startswith("Executive highlights") for x in tl) and tl[-1] != ""


def test_markup_and_long_text_are_plain(env, buckets):
    env.bm.put(f"/api/fundraisers/{buckets['id']}/summary-notes",
               {"summary_notes": "<script>alert(1)</script> & " + "very long note " * 60})
    r = deck(env, buckets)
    assert r.status_code == 200
    assert "<script>alert(1)</script>" in all_text(r.content)
    Presentation(io.BytesIO(r.content))      # opens


def test_cancelled_title_banner(env, buckets):
    env.bm.post(f"/api/fundraisers/{buckets['id']}/cancel", {"reason": "Storm warning"})
    s0 = slides(deck(env, buckets).content)[0]
    assert any("Cancelled: Storm warning" in sh.text_frame.text for sh in s0.shapes if sh.has_text_frame)


def test_access_and_audit(env, buckets):
    assert deck(env, buckets, client=env.admin).status_code == 403
    assert deck(env, buckets, client=env.auditor).status_code == 200
    ev = env.auditor.get("/api/audit-events?action=REPORT_GENERATED").json()["items"]
    assert any((e.get("after") or {}).get("report") == "FUNDRAISER_PRESENTATION" for e in ev)
    assert env.admin.put("/api/system/modules", {"fundraisers": False}).status_code == 200
    assert deck(env, buckets).status_code == 404
