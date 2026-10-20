"""2.1.0 (#277): fundraisers linked across years (series), entered prior-year results, "Create next year's from this
one", and the Fiscal Year fundraising report and presentation."""
from __future__ import annotations

import io
import zipfile

import pytest
from pptx import Presentation
from pypdf import PdfReader

from fmpoc.services import fundraiser_series as fseries


def pdf_text(content: bytes) -> str:
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(io.BytesIO(content)).pages)


@pytest.fixture
def years(env, base):
    """FY2027 = base; FY2026 earlier. Golf Outing in both years, a Raffle only this year."""
    assert env.admin.put("/api/system/modules", {"fundraisers": True}).status_code == 200
    f26 = env.fy("2026", "2025-07-01", "2026-06-30")
    inc26 = env.budget(f26["id"], "4100", "Golf", "INCOME", "1000.00")
    sel26 = {o["label"]: o["id"] for o in env.selectable(f26["id"], "DEPOSIT")}
    acct = base["acct"]["id"]
    env.txn(acct, "DEPOSIT", [{"budget_id": sel26["4100 Golf"], "amount": "800.00", "description": "Golf 2025"}], date="2025-09-20")
    g26 = env.bm.post("/api/fundraisers", {"name": "2025 Golf Outing", "start_date": "2025-09-20",
                                           "budget_ids": [inc26["id"]]}).json()
    golf27 = env.budget(base["fy"]["id"], "4100", "Golf", "INCOME", "1200.00")
    sel27 = {o["label"]: o["id"] for o in env.selectable(base["fy"]["id"], "DEPOSIT")}
    env.txn(acct, "DEPOSIT", [{"budget_id": sel27["4100 Golf"], "amount": "1000.00", "description": "Golf 2026"}], date="2026-09-19")
    g27 = env.bm.post("/api/fundraisers", {"name": "2026 Golf Outing", "start_date": "2026-09-19",
                                           "budget_ids": [golf27["id"]], "continues_from_id": g26["id"]})
    assert g27.status_code == 201, g27.text
    env.txn(acct, "DEPOSIT", [{"budget_id": base["inc_leaf"], "amount": "250.00", "description": "Raffle"}], date="2026-10-01")
    raffle = env.bm.post("/api/fundraisers", {"name": "Raffle", "start_date": "2026-10-01", "budget_ids": [base["inc"]["id"]],
                                              "filter_text": "Raffle"}).json()
    return {"f26": f26, "g26": g26, "g27": g27.json(), "raffle": raffle, "base": base}


def test_names_and_years():
    assert fseries.series_name_for("2026 Golf Outing") == "Golf Outing"
    assert fseries.series_name_for("Pancake Breakfast - 2025") == "Pancake Breakfast"
    assert fseries.next_name("Golf Outing 2026") == "Golf Outing 2027"
    assert fseries.label_year("FY2019") == 2019 and fseries.label_year("2018-19") == 2018 and fseries.label_year("Spring") is None


def test_continues_from_creates_a_series_and_joins_it(env, years):
    g27 = years["g27"]
    assert g27["series_link"]["series"]["name"] == "Golf Outing"
    assert g27["series_link"]["continues_from"]["id"] == years["g26"]["id"]
    s = env.bm.get("/api/fundraiser-series").json()
    assert [x["name"] for x in s] == ["Golf Outing"]
    assert [f["name"] for f in s[0]["fundraisers"]] == ["2025 Golf Outing", "2026 Golf Outing"]
    acts = {e["action"] for e in env.auditor.get("/api/audit-events").json()["items"]}
    assert "FUNDRAISER_SERIES_CREATED" in acts


def test_continue_options_latest_per_series_and_unlinked(env, years):
    # for a 2028 fundraiser: the latest Golf Outing (not the 2025 one) and the unlinked Raffle, newest first
    f28 = env.fy("2028", "2027-07-01", "2028-06-30")
    opts = env.bm.get("/api/fundraisers/continue-options?start_date=2027-09-18").json()
    assert [o["name"] for o in opts] == ["Raffle", "2026 Golf Outing"]
    # a skipped year still links: from 2029 the same options
    env.fy("2029", "2028-07-01", "2029-06-30")
    assert [o["name"] for o in env.bm.get("/api/fundraisers/continue-options?start_date=2028-09-18").json()] == ["Raffle", "2026 Golf Outing"]
    # same-year fundraisers are not offered
    assert [o["name"] for o in env.bm.get("/api/fundraisers/continue-options?start_date=2026-11-01").json()] == ["2025 Golf Outing"]
    assert f28


def test_link_changes_in_a_closed_year_are_audited(env, years, app):
    from sqlalchemy.orm import Session
    from fmpoc.models import FiscalYear
    with Session(app.state.engine) as db:
        db.get(FiscalYear, years["f26"]["id"]).status = "CLOSED"
        db.get(FiscalYear, years["base"]["fy"]["id"]).status = "CLOSED"
        db.commit()
    gid = years["g27"]["id"]
    r = env.bm.put(f"/api/fundraisers/{gid}/series-link", {"continues_from_id": None})
    assert r.status_code == 200 and r.json()["series_link"]["series"] is None
    r = env.bm.put(f"/api/fundraisers/{gid}/series-link", {"continues_from_id": years["g26"]["id"]})
    assert r.status_code == 200 and r.json()["series_link"]["series"]["name"] == "Golf Outing"
    ev = env.auditor.get(f"/api/audit-events?action=FUNDRAISER_SERIES_LINK_UPDATED").json()["items"]
    assert len(ev) == 2
    # a later fundraiser can't be "continued from"
    r = env.bm.put(f"/api/fundraisers/{years['g26']['id']}/series-link", {"continues_from_id": gid})
    assert r.status_code == 422


def test_entered_results_and_merge(env, years):
    sid = env.bm.get("/api/fundraiser-series").json()[0]["id"]
    r = env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2024", "net": "650.00",
                                                              "note": "estimated, from the old spreadsheet"})
    assert r.status_code == 201, r.text
    # refused: a year with a fundraiser in PennyWarden, a duplicate, no year in the label
    assert env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2026", "net": "1.00"}).status_code == 409
    assert env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2024", "net": "1.00"}).status_code == 409
    assert env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "Spring", "net": "1.00"}).status_code == 422
    rid = r.json()[0]["results"][0]["id"]
    # a second series by mistake, with an entered result for the same year: merge refused until one is retired
    other_f = env.bm.post("/api/fundraisers", {"name": "Golf (old name) 2022", "start_date": "2022-09-10"}).json()
    env.bm.put(f"/api/fundraisers/{years['raffle']['id']}/series-link", {"continues_from_id": other_f["id"]})
    lst = {x["name"]: x for x in env.bm.get("/api/fundraiser-series").json()}
    s2 = lst["Golf (old name)"]["id"]
    env.bm.post(f"/api/fundraiser-series/{s2}/results", {"fiscal_year_label": "FY2024", "net": "10.00"})
    r = env.bm.post(f"/api/fundraiser-series/{s2}/merge", {"into_series_id": sid})
    assert r.status_code == 409 and r.json()["error"]["code"] == "ENTERED_RESULT_CONFLICT"
    s2_rid = next(x for x in env.bm.get("/api/fundraiser-series").json() if x["id"] == s2)["results"][0]["id"]
    assert env.bm.post(f"/api/fundraiser-series/{s2}/results/{s2_rid}/retire", {"reason": "duplicate"}).status_code == 200
    r = env.bm.post(f"/api/fundraiser-series/{s2}/merge", {"into_series_id": sid})
    assert r.status_code == 200, r.text
    golf = next(x for x in r.json() if x["id"] == sid)
    assert {f["name"] for f in golf["fundraisers"]} >= {"Raffle", "Golf (old name) 2022"}
    assert any(x["retired"] for x in golf["results"]) and any(x["id"] == rid and not x["retired"] for x in golf["results"])
    # rename and Group as Other
    r = env.bm.patch(f"/api/fundraiser-series/{sid}", {"name": "Golf Outing & Raffle", "group_as_other": True})
    assert r.status_code == 200 and r.json()[0]["group_as_other"] is True
    acts = {e["action"] for e in env.auditor.get("/api/audit-events").json()["items"]}
    assert {"FUNDRAISER_RESULT_ENTERED", "FUNDRAISER_RESULT_RETIRED", "FUNDRAISER_SERIES_MERGED", "FUNDRAISER_SERIES_UPDATED"} <= acts


def test_next_year_draft_and_bucket_copy(env, years):
    gid = years["g27"]["id"]
    env.ru.post(f"/api/fundraisers/{gid}/categories", {"name": "Day-of"})
    env.ru.post(f"/api/fundraisers/{gid}/buckets", {"name": "Hole sponsors"})
    f28 = env.fy("2028", "2027-07-01", "2028-06-30")
    golf28 = env.budget(f28["id"], "4100", "Golf", "INCOME", "1300.00")
    env.bm.patch(f"/api/budgets/{golf28['id']}", {"continues_budget_id": years["g27"]["budgets"][0]["id"]})
    d = env.bm.get(f"/api/fundraisers/{gid}/next-year").json()
    assert d["name"] == "2027 Golf Outing" and d["start_date"] == "2027-09-19" and d["continues_from_id"] == gid
    assert d["budget_ids"] == [golf28["id"]] and d["buckets"] == ["Hole sponsors"] and d["categories"] == ["Day-of"]
    body = {k: d[k] for k in ("name", "description", "start_date", "end_date", "filter_text", "filter_regex", "budget_ids",
                              "continues_from_id", "copy_buckets_from_id")}
    r = env.bm.post("/api/fundraisers", body)
    assert r.status_code == 201, r.text
    n = r.json()
    assert [b["name"] for b in n["buckets"]["items"]] == ["Hole sponsors"] and n["buckets"]["categories"][0]["name"] == "Day-of"
    assert n["series_link"]["series"]["name"] == "Golf Outing"


def _report(env, ys, client=None, fmt="json", **q):
    params = {"fiscal_year_id": ys["base"]["fy"]["id"], "ids": [ys["g27"]["id"], ys["raffle"]["id"]], "format": fmt,
              **{k: str(v).lower() if isinstance(v, bool) else v for k, v in q.items()}}
    return (client or env.bu).get("/api/fundraisers/fy-report", params=params)


def test_report_figures_history_and_defaults(env, years):
    sel = env.bu.get(f"/api/fundraisers/fy-report/selection?fiscal_year_id={years['base']['fy']['id']}").json()
    assert sorted(sel["default_ids"]) == sorted([years["g27"]["id"], years["raffle"]["id"]]) and sel["year"] == 2027
    rep = _report(env, years).json()
    assert rep["year_total"] == "1250.00"
    rows = {r["name"]: r for r in rep["this_year"]}
    page = env.bu.get(f"/api/fundraisers/{years['g27']['id']}").json()["totals"]
    assert rows["2026 Golf Outing"]["net_proceeds"] == page["net"] == "1000.00"
    assert rows["2026 Golf Outing"]["share"] == "80.0"
    assert [(t["year"], t["total"]) for t in rep["trend"]] == [(2026, "800.00"), (2027, "1250.00")]
    assert rep["growth_vs_last"] == "56.3" and rep["peak_year"] == 2027 and rep["cumulative"] == "2050.00"
    golf = next(a for a in rep["averages"] if a["name"] == "Golf Outing")
    assert golf["active_years"] == 2 and golf["average"] == "900.00"
    assert ("Golf Outing", 2026, "Recorded") in {(h["series"], h["year"], h["source"]) for h in rep["history"]}
    # Raffle is under 5%? no: 250 of 2050 = 12.2% - its own slice
    assert {x["name"] for x in rep["share"]} == {"Golf Outing", "Raffle"}


def test_two_fundraisers_of_one_series_in_a_year_add_together(env, years):
    fall = env.bm.post("/api/fundraisers", {"name": "2026 Golf Outing (fall)", "start_date": "2026-11-01",
                                            "continues_from_id": years["g26"]["id"]}).json()
    rep = _report(env, years, ids=[years["g27"]["id"], years["raffle"]["id"], fall["id"]]).json()
    golf27 = [h for h in rep["history"] if h["series"] == "Golf Outing" and h["year"] == 2027]
    assert len(golf27) == 1 and golf27[0]["amount"] == "1000.00"
    assert len(rep["this_year"]) == 3


def test_entered_years_targeted_total_and_other(env, years):
    sid = env.bm.get("/api/fundraiser-series").json()[0]["id"]
    env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2025", "net": "700.00"})
    pancake = env.bm.post("/api/fundraisers", {"name": "2025 Pancake Breakfast", "start_date": "2025-11-01"}).json()
    env.bm.post("/api/fundraisers", {"name": "2026 Pancake Breakfast", "start_date": "2027-03-01",
                                     "continues_from_id": pancake["id"]})
    psid = next(x["id"] for x in env.bm.get("/api/fundraiser-series").json() if x["name"] == "Pancake Breakfast")
    env.bm.post(f"/api/fundraiser-series/{psid}/results", {"fiscal_year_label": "FY2024", "net": "40.00"})
    rep = _report(env, years, years=-1, targeted=True).json()
    assert [t["year"] for t in rep["trend"]] == [2024, 2025, 2026, 2027]
    assert ("Golf Outing", 2025, "Entered") in {(h["series"], h["year"], h["source"]) for h in rep["history"]}
    # Pancake Breakfast has not been held this year: its average of earlier years is the estimate
    est = {e["series"]: e for e in rep["estimates"]}
    assert est["Pancake Breakfast"]["estimate"] == "20.00"     # (40 entered for 2024 + 0 recorded for 2026) / 2
    assert rep["targeted_total"] == "1270.00"
    other = next(x for x in rep["share"] if x["name"] == "Other")
    assert "Pancake Breakfast" in other["members"]                       # under 5% of the period
    assert any("Targeted total" in x for x in rep["footnotes"]) and any("Entered" in x for x in rep["footnotes"])


def test_pdf_and_pptx(env, years):
    env.bm.put("/api/fundraisers/fy-report/notes", {"fiscal_year_id": years["base"]["fy"]["id"],
                                                     "notes": "A matching-gift program is expected to add revenue."})
    r = _report(env, years, fmt="pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    t = pdf_text(r.content)
    for needle in ("Fundraising — FY2027", "Executive highlights", "This year by fundraiser", "2026 Golf Outing", "Raffle",
                   "$1,250.00", "Annual trend", "Share by series", "Average per active year", "Historical results",
                   "Recorded", "matching-gift program", "Page 1 of"):
        assert needle in t, needle
    r = _report(env, years, fmt="pptx")
    assert r.status_code == 200
    prs = Presentation(io.BytesIO(r.content))
    assert all(sl.notes_slide.notes_text_frame.text.strip() for sl in prs.slides)
    kinds = {sh.chart.chart_type for sl in prs.slides for sh in sl.shapes if sh.has_chart}
    from pptx.enum.chart import XL_CHART_TYPE
    assert {XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.BAR_CLUSTERED} <= kinds
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert not any(n.startswith("ppt/media/") or n.endswith(".bin") for n in z.namelist())
    # notes remembered per Fiscal Year; audited
    assert "matching-gift" in env.bu.get(f"/api/fundraisers/fy-report/selection?fiscal_year_id={years['base']['fy']['id']}").json()["notes"]
    ev = [e for e in env.auditor.get("/api/audit-events?action=REPORT_GENERATED").json()["items"]
          if (e.get("after") or {}).get("report", "").startswith("FUNDRAISER_FY")]
    assert {e["after"]["report"] for e in ev} == {"FUNDRAISER_FY_SUMMARY", "FUNDRAISER_FY_PRESENTATION"}
    assert ev[0]["after"]["fundraiser_ids"]


def test_no_linked_history(env, base):
    assert env.admin.put("/api/system/modules", {"fundraisers": True}).status_code == 200
    f = env.bm.post("/api/fundraisers", {"name": "Solo", "start_date": "2026-09-12", "budget_ids": [base["inc"]["id"]]}).json()
    q = {"fiscal_year_id": base["fy"]["id"], "ids": [f["id"]]}
    rep = env.bu.get("/api/fundraisers/fy-report", params={**q, "format": "json"}).json()
    assert rep["linked_history"] is False
    assert "No earlier years linked yet." in pdf_text(env.bu.get("/api/fundraisers/fy-report", params={**q, "format": "pdf"}).content)
    assert env.bu.get("/api/fundraisers/fy-report", params={**q, "format": "pptx"}).status_code == 200


def test_permissions(env, years):
    sid = env.bm.get("/api/fundraiser-series").json()[0]["id"]
    for c in (env.bu, env.ru, env.auditor, env.admin):
        assert c.get("/api/fundraiser-series").status_code == 403
        assert c.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2020", "net": "1.00"}).status_code == 403
        assert c.put(f"/api/fundraisers/{years['g27']['id']}/series-link", {"continues_from_id": None}).status_code == 403
        assert c.put("/api/fundraisers/fy-report/notes", {"fiscal_year_id": years["base"]["fy"]["id"], "notes": "x"}).status_code == 403
    assert _report(env, years, client=env.admin).status_code == 403
    assert _report(env, years, client=env.auditor).status_code == 200
    assert env.admin.put("/api/system/modules", {"fundraisers": False}).status_code == 200
    assert _report(env, years).status_code == 404


def test_next_year_before_its_fiscal_year_exists(env, years):
    """Planned before next year's Fiscal Year is set up: still "next year" (2028), so this year's fundraisers are offered."""
    opts = env.bm.get("/api/fundraisers/continue-options?start_date=2027-09-18").json()
    assert {o["name"] for o in opts} == {"Raffle", "2026 Golf Outing"}
    assert fseries.date_year.__doc__


# ------------------------------------------------------------------ year over year (Kyle, 2026-10-20)
def test_year_over_year_in_the_fundraiser_report_and_deck(env, years):
    sid = env.bm.get("/api/fundraiser-series").json()[0]["id"]
    env.bm.post(f"/api/fundraiser-series/{sid}/results", {"fiscal_year_label": "FY2025", "net": "600.00"})
    from sqlalchemy.orm import Session
    from fmpoc.models import Fundraiser, User
    from fmpoc.services import fundraiser_summary as fsum
    with Session(env.app.state.engine) as db:
        f = db.get(Fundraiser, years["g27"]["id"])

        class Ctx:
            workspace_id = f.workspace_id
            user = db.query(User).filter_by(username="budgetuser").one()
        s = fsum.summary(db, Ctx, f)
    yy = s["year_over_year"]
    assert [(r["year"], r["source"], r["net_proceeds"]) for r in yy] == [(2025, "Entered", "600.00"), (2026, "Recorded", "800.00"),
                                                                        (2027, "Recorded", "1000.00")]
    assert yy[0]["net_revenue"] is None                          # entered with the net only
    assert (yy[2]["change"], yy[2]["change_pct"]) == ("200.00", "25.0")
    assert any("Net proceeds were up $200.00 (25.0%) on 2026" in x for x in s["takeaways"])
    t = pdf_text(env.bu.get(f"/api/fundraisers/{years['g27']['id']}/summary-report").content)
    assert "Year over year" in t and "up $200.00 (25.0%)" in t and "FY2025 (Entered)" in t
    prs = Presentation(io.BytesIO(env.bu.get(f"/api/fundraisers/{years['g27']['id']}/presentation").content))
    titles = [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text), "") for sl in prs.slides]
    assert "Year over year" in titles
    # last year's report goes up to its own year only (the entered 2025, then itself), never later years
    t0 = pdf_text(env.bu.get(f"/api/fundraisers/{years['g26']['id']}/summary-report").content)
    assert "Year over year" in t0 and "up $200.00 (33.3%)" in t0 and "2026 Golf Outing" not in t0
    # and a fundraiser in no series neither
    assert "Year over year" not in pdf_text(env.bu.get(f"/api/fundraisers/{years['raffle']['id']}/summary-report").content)


def test_year_over_year_by_series_in_the_fiscal_year_report(env, years):
    rep = _report(env, years).json()
    golf = next(x for x in rep["year_over_year"] if x["name"] == "Golf Outing")
    assert golf["values"] == ["800.00", "1000.00"] and golf["change"] == "200.00" and golf["change_pct"] == "25.0"
    raffle = next(x for x in rep["year_over_year"] if x["name"] == "Raffle")
    assert raffle["values"] == [None, "250.00"] and raffle["change"] is None
    t = pdf_text(_report(env, years, fmt="pdf").content)
    assert "Year over year by series" in t and "up $200.00 (25.0%)" in t
    prs = Presentation(io.BytesIO(_report(env, years, fmt="pptx").content))
    titles = [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text), "") for sl in prs.slides]
    assert "Year over year by series" in titles and "Change on the year before" in titles
