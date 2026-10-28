"""2.2.0 (#280): organization branding for generated reports and presentations - Material 3 palettes (ready-made, or
custom resolved colors), WCAG AA contrast refused, similar series colors confirmed, logos validated and re-encoded,
applied to PDF reports and decks, in backups as its own restore section and kept through a reinitialize."""
from __future__ import annotations

import io
import json

import pytest
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from conftest import PASSWORD
from test_v141_backup import _backup, _restore
from test_v161_fundraiser_manage import gala, put_line  # noqa: F401  (fixture)
from test_v210_fundraiser_summary import buckets  # noqa: F401  (fixture)

from fmpoc.services import branding as br
from fmpoc.services import partial_restore as pr

URL = "/api/system/branding"


def _png(w=400, h=100, color=(200, 30, 30), fmt="PNG", meta=False) -> bytes:
    im = Image.new("RGB", (w, h), color)
    out = io.BytesIO()
    kw = {}
    if meta and fmt == "PNG":
        from PIL.PngImagePlugin import PngInfo
        info = PngInfo()
        info.add_text("Author", "SECRET-METADATA")
        kw["pnginfo"] = info
    im.save(out, fmt, **kw)
    return out.getvalue()


def _upload(client, slot, data, name="logo.png", ctype="image/png"):
    return client.c.post(f"{URL}/logo/{slot}", files={"file": (name, data, ctype)},
                         headers={"X-CSRF-Token": client.csrf})


def _custom(base_preset="ocean", **over):
    p = br.PRESETS[base_preset]
    body = {"preset": "custom", "source_primary": p["source_primary"], "source_secondary": p["source_secondary"],
            "colors": dict(p["colors"]), "series": list(p["series"]), "footer_text": None}
    body.update(over)
    return body


def test_presets_meet_contrast_and_default_is_pennywarden(env, base):
    for k, p in br.PRESETS.items():
        assert all(x["ok"] for x in br.contrast_report(p["colors"])), k
        assert len(p["series"]) == br.SERIES_COUNT and not br.similar_series(p["series"]), k
    v = env.admin.get(URL).json()
    assert v["preset"] == "pennywarden" and v["logos"] == {"full": None, "mark": None} and v["footer_text"] is None
    assert [p["id"] for p in v["presets"]] == ["pennywarden", "forest", "ocean", "plum", "ember"]


def test_only_administrators(env, base):
    for c in (env.bm, env.bu, env.auditor):
        assert c.get(URL).status_code == 403
        assert c.put(URL, {"preset": "forest"}).status_code == 403
        assert _upload(c, "full", _png()).status_code == 403


def test_preset_custom_contrast_similar_and_reset(env, base):
    adm = env.admin
    r = adm.put(URL, {"preset": "forest", "footer_text": "Serving the community since 1952"})
    assert r.status_code == 200, r.text
    assert r.json()["colors"]["primary"] == br.PRESETS["forest"]["colors"]["primary"]
    # custom: a text color that fails WCAG AA is refused with the reason
    bad = _custom(colors={**br.PRESETS["ocean"]["colors"], "on_primary": "#3a7fa8"})
    r = adm.put(URL, bad)
    assert r.status_code == 422 and "4.5:1" in r.text and "WCAG AA" in r.text
    # similar series colors need a confirmation
    s = list(br.PRESETS["ocean"]["series"])
    s[1] = "#01668f"
    r = adm.put(URL, _custom(series=s))
    assert r.status_code == 409 and r.json()["error"]["warnings"][0]["code"] == "SIMILAR_SERIES_COLORS"
    r = adm.put(URL, {**_custom(series=s), "footer_text": "Serving the community since 1952",
                      "confirmations": ["SIMILAR_SERIES_COLORS"]})
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["preset"] == "custom" and v["series"][1] == "#01668f" and v["similar_series"] == [[0, 1]]
    # validation: hex, every role, eight series, one footer line
    assert adm.put(URL, _custom(source_primary="blue")).status_code == 422
    c = dict(br.PRESETS["ocean"]["colors"])
    del c["outline"]
    assert adm.put(URL, _custom(colors=c)).status_code == 422
    assert adm.put(URL, _custom(series=s[:5])).status_code == 422
    assert adm.put(URL, {"preset": "forest", "footer_text": "x" * 121}).status_code == 422
    assert adm.put(URL, {"preset": "forest", "footer_text": "a\nb"}).status_code == 422
    assert adm.put(URL, {"preset": "neon"}).status_code == 422
    # reset: the palette only
    r = adm.post(f"{URL}/reset")
    assert r.json()["preset"] == "pennywarden" and r.json()["footer_text"] == "Serving the community since 1952"
    ev = [e for e in env.auditor.get("/api/audit-events").json()["items"] if e["action"] == "BRANDING_UPDATED"]
    assert len(ev) >= 3


def test_logos_validated_reencoded_and_audited(env, base):
    adm = env.admin
    r = _upload(adm, "full", _png(meta=True))
    assert r.status_code == 200, r.text
    info = r.json()["logos"]["full"]
    assert (info["width"], info["height"]) == (400, 100)
    img = adm.get(f"{URL}/logo/full")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png"
    assert b"SECRET-METADATA" not in img.content                     # re-encoded: metadata stripped
    assert _upload(adm, "mark", _png(64, 64, fmt="JPEG"), "m.jpg", "image/jpeg").status_code == 200
    assert _upload(adm, "full", b"<svg xmlns='http://www.w3.org/2000/svg'/>", "x.svg", "image/svg+xml").status_code == 415
    assert _upload(adm, "full", b"\x89PNG\r\n\x1a\n" + b"x" * 100).status_code == 415
    assert _upload(adm, "full", b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024 + 10)).status_code == 413
    assert _upload(adm, "banner", _png()).status_code == 422
    assert adm.post(f"{URL}/logo/mark/remove").json()["logos"]["mark"] is None
    assert adm.get(f"{URL}/logo/mark").status_code == 404
    ev = [e for e in env.auditor.get("/api/audit-events?page_size=200").json()["items"]
          if e["action"] in ("BRANDING_LOGO_SET", "BRANDING_LOGO_REMOVED")]
    assert len(ev) == 3 and "SECRET" not in json.dumps(ev)


def _flow_pdf(env, base) -> PdfReader:
    body = {"title": "Branding check", "date_from": "2026-08-01", "date_to": "2026-08-31",
            "bank_account_ids": [base["acct"]["id"]]}
    r = env.bu.post("/api/reports/financial-flow", body)
    assert r.status_code == 200, r.text
    return PdfReader(io.BytesIO(r.content))


def test_pdf_reports_use_the_branding(env, base):
    plain = _flow_pdf(env, base)
    assert len(plain.pages[0].images) == 0
    adm = env.admin
    assert adm.put(URL, {**_custom(), "footer_text": "Acme Org - Building Hope"}).status_code == 200
    assert _upload(adm, "full", _png()).status_code == 200
    assert _upload(adm, "mark", _png(64, 64)).status_code == 200
    pdf = _flow_pdf(env, base)
    assert len(pdf.pages[0].images) == 2                     # the full logo (first page) and the mark
    if len(pdf.pages) > 1:
        assert len(pdf.pages[1].images) == 1
    assert "Acme Org - Building Hope" in pdf.pages[0].extract_text()
    # the headings are in the palette's primary color
    r, g, b = (int(br.PRESETS["ocean"]["colors"]["primary"][i:i + 2], 16) / 255 for i in (1, 3, 5))
    from reportlab.lib.rl_accel import fp_str
    raw = pdf.pages[0].get_contents().get_data().decode("latin-1")
    assert f"{fp_str(r, g, b)} rg" in raw


def test_deck_uses_the_branding(env, base, buckets):  # noqa: F811
    g = buckets
    adm = env.admin
    assert adm.put(URL, {"preset": "plum", "footer_text": "Plum footer"}).status_code == 200
    assert _upload(adm, "full", _png()).status_code == 200
    assert _upload(adm, "mark", _png(64, 64)).status_code == 200
    r = env.bu.get(f"/api/fundraisers/{g['id']}/presentation")
    assert r.status_code == 200, r.text
    prs = Presentation(io.BytesIO(r.content))
    sl = list(prs.slides)
    pics = [sh for sh in sl[0].shapes if sh.shape_type == 13]
    assert len(pics) == 1                                   # the full logo instead of the coin
    later = [sh for sh in sl[1].shapes if sh.shape_type == 13]
    assert len(later) == 1 and any(sh.has_text_frame and sh.text_frame.text == "Plum footer" for sh in sl[1].shapes)
    xml = b"".join(p.blob for p in prs.part.package.iter_parts() if p.partname.endswith(".xml"))
    assert br.PRESETS["plum"]["series"][0][1:].upper().encode() in xml
    # the stakeholder report (PDF and Word) too
    assert env.bu.get(f"/api/fundraisers/{g['id']}/summary-report").status_code == 200
    d = env.bu.get(f"/api/fundraisers/{g['id']}/summary-report?format=docx")
    assert d.status_code == 200
    import zipfile
    doc = zipfile.ZipFile(io.BytesIO(d.content)).read("word/document.xml")
    assert br.PRESETS["plum"]["colors"]["primary"][1:].upper().encode() in doc


def test_restore_section_and_reinitialize(env, base):
    assert "settings.branding" in pr.DOMAINS
    adm = env.admin
    assert adm.put(URL, {"preset": "ember"}).status_code == 200
    assert _upload(adm, "full", _png()).status_code == 200
    blob = _backup(env)
    assert adm.put(URL, {"preset": "forest"}).status_code == 200
    adm.post(f"{URL}/logo/full/remove")
    j = _restore(adm, blob, password=PASSWORD, confirm="RESTORE", domains=["settings.general"])
    assert j["state"] == "done", j
    from conftest import Api
    a2 = Api(env.app)
    assert a2.login("admin").status_code == 200
    assert a2.get(URL).json()["preset"] == "forest"                 # Branding not chosen: kept
    j = _restore(a2, blob, password=PASSWORD, confirm="RESTORE", domains=["settings.branding"])
    assert j["state"] == "done", j
    a3 = Api(env.app)
    assert a3.login("admin").status_code == 200
    v = a3.get(URL).json()
    assert v["preset"] == "ember" and v["logos"]["full"] is not None
    from fmpoc.models import BrandingSettings
    from fmpoc.services.reinitialize import WIPE
    assert BrandingSettings not in WIPE                              # kept through a reinitialize


@pytest.mark.parametrize("bad", ["#12345", "123456", "#gggggg"])
def test_hex_values(env, base, bad):
    assert env.admin.put(URL, _custom(colors={**br.PRESETS["ocean"]["colors"], "primary": bad})).status_code == 422
