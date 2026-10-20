"""2.1.0 (#275): the fundraiser stakeholder report (Financial & Proceeds Statement) as PDF and Word, its options,
the status line (Interim / Post-event / Final / Cancelled), summary notes and Mark results final - and that the audit
detail is unchanged."""
from __future__ import annotations

import io
import zipfile

import docx
import pytest
from pypdf import PdfReader

from fmpoc.services import fundraiser_summary as fs
from test_v161_fundraiser_manage import gala, put_line  # noqa: F401  (fixture)


def url(g, tail=""):
    return f"/api/fundraisers/{g['id']}{tail}"


def pdf_text(content: bytes) -> str:
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(io.BytesIO(content)).pages)


def docx_text(content: bytes) -> str:
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    for sec in d.sections:
        parts += [p.text for p in sec.footer.paragraphs]
    return " ".join(" ".join(parts).split())


@pytest.fixture
def buckets(env, gala):
    """Float 200 out / 200 returned; Food Sales (mixed: 1000 sales, 300 supplies as charge 450 less 150 credit),
    Raffle 250 (in Day-of event); the 75 donation unassigned (Other)."""
    g = gala
    r = lambda name, **kw: env.ru.post(url(g, "/buckets"), {"name": name, **kw}).json()
    cat = env.ru.post(url(g, "/categories"), {"name": "Day-of event"}).json()["buckets"]["categories"][0]
    d = r("Food Sales", category_id=cat["id"])
    d = r("Raffle", category_id=cat["id"])
    ids = {b["name"]: b["id"] for b in d["buckets"]["items"]}
    put_line(env.ru, g, "Cash float", {"classification": {"kind": "CASH_FLOAT_OUT", "amount": "200.00"}})
    put_line(env.ru, g, "Gala takings", {"classification": {"kind": "CASH_FLOAT_RETURNED", "amount": "200.00"},
                                         "buckets": [{"bucket_id": ids["Food Sales"], "amount": "1000.00"},
                                                     {"bucket_id": ids["Raffle"], "amount": "250.00"}]})
    put_line(env.ru, g, "Food supplies", {"buckets": [{"bucket_id": ids["Food Sales"], "charge": "450.00", "less": "150.00",
                                                       "less_description": "Pre-paid team credit"}]})
    return {**g, "ids": ids}


def get(env, g, client=None, **q):
    params = {k: str(v).lower() if isinstance(v, bool) else v for k, v in q.items()}
    return (client or env.bu).get(url(g, "/summary-report"), params=params)


def test_figures_match_the_fundraiser_page(env, buckets):
    page = env.bu.get(url(buckets)).json()["totals"]
    from sqlalchemy.orm import Session
    from fmpoc.models import Fundraiser, User
    with Session(env.app.state.engine) as db:
        f = db.get(Fundraiser, buckets["id"])

        class Ctx:
            workspace_id = f.workspace_id
            user = db.query(User).filter_by(username="budgetuser").one()
        s = fs.summary(db, Ctx, f)
    t = s["totals"]
    assert (t["net_revenue"], t["expenses"], t["net_proceeds"]) == (page["income"], page["expense"], page["net"])
    assert (t["net_revenue"], t["expenses"], t["net_proceeds"]) == ("1325.00", "300.00", "1025.00")
    assert t["gross_deposits"] == "1525.00" and t["cash_float_returned"] == "200.00" and t["cash_float_out"] == "200.00"
    assert t["net_margin"] == "77.4" and t["roi"] == page["roi"] and t["credits"] == "150.00"
    # breakdowns add up to the headline totals, Other included
    inc = sum(float(g_["income"]) for g_ in s["groups"]) + float(s["other"]["income"])
    exp = sum(float(g_["expense"]) for g_ in s["groups"]) + float(s["other"]["expense"])
    assert round(inc, 2) == 1325.00 and round(exp, 2) == 300.00
    day = s["groups"][0]
    assert day["name"] == "Day-of event" and [b["name"] for b in day["buckets"]] == ["Food Sales", "Raffle"]
    food = day["buckets"][0]
    assert (food["income"], food["expense"], food["net"]) == ("1000.00", "300.00", "700.00")   # mixed bucket
    assert s["other"]["income"] == "75.00"
    x = food["expense_lines"][0]
    assert (x["charge"], x["less"], x["amount"], x["less_description"]) == ("450.00", "150.00", "300.00", "Pre-paid team credit")
    texts = " ".join(s["takeaways"])
    assert "net margin of 77.4%" in texts and "Day-of event was the largest source of revenue" in texts
    assert "Cash float: $200.00 taken out, $200.00 returned (fully returned)" in texts
    assert "Credits received on expenses: $150.00" in texts
    assert "Day-of event income alone covered all expenses" in texts


def test_pdf_sections_and_nothing_from_the_audit_detail(env, buckets):
    r = get(env, buckets)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    t = pdf_text(r.content)
    for needle in ("Financial & Proceeds Statement", "Gala", "Executive summary", "Total net revenue", "$1,325.00",
                   "Net proceeds", "$1,025.00", "Net margin 77.4%", "Financial summary", "Gross deposits", "$1,525.00",
                   "Less cash float returned", "ROI", "Revenue breakdown", "Day-of event", "Food Sales",
                   "Other (not in a bucket)", "Itemized revenue", "Expense breakdown", "Food supplies — charge",
                   "Less: Pre-paid team credit", "-$150.00", "Net by category / bucket", "Key takeaways",
                   "Prepared for committee & stakeholders", "Transaction-level evidence is in the Audit report", "Page 1 of"):
        assert needle in t, needle
    # no budgets, accounts, ids or excluded-line details
    for banned in ("Operations", "Donations", "Acct ", "1000 ", "4000 ", "Budget", "Excluded"):
        assert banned not in t, banned
    assert len(PdfReader(io.BytesIO(r.content)).pages[0].images) == 0


def test_docx_same_sections_and_figures(env, buckets):
    r = get(env, buckets, format="docx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert r.headers["content-disposition"].startswith('attachment; filename="fundraiser-Gala-report.docx"')
    t = docx_text(r.content)
    for needle in ("Financial & Proceeds Statement", "Executive summary", "$1,325.00", "$1,025.00", "Net margin 77.4%",
                   "Revenue breakdown", "Itemized revenue", "Expense breakdown", "Less: Pre-paid team credit",
                   "Net by category / bucket", "Key takeaways", "Prepared for committee & stakeholders"):
        assert needle in t, needle
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert not any(n.startswith("word/media/") for n in names)          # no pictures / attachments
    assert not any("vbaProject" in n for n in names)                     # no macros
    footer = " ".join(z.read(n).decode() for n in names if n.startswith("word/footer"))
    assert "PAGE" in footer and "NUMPAGES" in footer


def test_options(env, buckets):
    off = pdf_text(get(env, buckets, itemized=False).content)
    assert "Itemized revenue" not in off and "Revenue breakdown" in off
    # contributor names off: no Entity names anywhere (PDF and Word)
    ent = env.entity("Jordan Benefactor")
    env.txn(buckets["base"]["acct"]["id"], "DEPOSIT", [{"budget_id": buckets["base"]["inc_leaf"], "amount": "40.00",
                                                       "description": "Sponsor"}], date="2026-09-15", entity_id=ent["id"])
    on = pdf_text(get(env, buckets).content)
    assert "Jordan Benefactor" in on
    for fmt_ in ("pdf", "docx"):
        r = get(env, buckets, names=False, format=fmt_)
        t = pdf_text(r.content) if fmt_ == "pdf" else docx_text(r.content)
        assert "Jordan Benefactor" not in t
        assert "Contributor 1" in t


def test_no_buckets_and_plain_line_list(env, gala):
    t = pdf_text(get(env, gala).content)
    assert "Revenue breakdown" not in t and "Expense breakdown" not in t and "Contributor / description" not in t
    t = pdf_text(get(env, gala, plain_lines=True).content)
    assert "Contributor / description" in t and "Gala takings" in t
    d = docx_text(get(env, gala, plain_lines=True, format="docx").content)
    assert "Contributor / description" in d


def test_status_line_final_and_notes(env, buckets):
    assert "· Post-event" in pdf_text(get(env, buckets).content)        # the event (9/12/2026) is over
    st = env.bu.get(url(buckets)).json()["report_status"]
    assert st["code"] in ("INTERIM", "POST_EVENT") and st["final"] is False
    # Budget Managers only
    for c in (env.ru, env.bu, env.auditor):
        assert c.put(url(buckets, "/results-final"), {"final": True}).status_code == 403
        assert c.put(url(buckets, "/summary-notes"), {"summary_notes": "x"}).status_code == 403
    r = env.bm.put(url(buckets, "/results-final"), {"final": True})
    assert r.status_code == 200 and r.json()["report_status"]["code"] == "FINAL"
    assert "Final" in pdf_text(get(env, buckets).content)
    r = env.bm.put(url(buckets, "/results-final"), {"final": False})      # reversible
    assert r.json()["report_status"]["final"] is False
    r = env.bm.put(url(buckets, "/summary-notes"), {"summary_notes": "Thanks to <b>all</b> 40 volunteers."})
    assert r.status_code == 200 and r.json()["summary_notes"] == "Thanks to <b>all</b> 40 volunteers."
    t = pdf_text(get(env, buckets).content)
    assert "Summary notes" in t and "Thanks to <b>all</b> 40 volunteers." in t        # escaped, printed as text
    assert "Thanks to <b>all</b> 40 volunteers." in docx_text(get(env, buckets, format="docx").content)
    acts = [e["action"] for e in env.auditor.get(f"/api/audit-events?object_type=fundraiser&object_id={buckets['id']}").json()["items"]]
    assert {"FUNDRAISER_RESULTS_FINAL_MARKED", "FUNDRAISER_RESULTS_FINAL_UNMARKED", "FUNDRAISER_SUMMARY_NOTES_UPDATED"} <= set(acts)
    gen = [e for e in env.auditor.get("/api/audit-events?action=REPORT_GENERATED").json()["items"]
           if (e.get("after") or {}).get("report") == "FUNDRAISER_SUMMARY"]
    assert {e["after"]["format"] for e in gen} == {"PDF", "DOCX"}


def test_final_when_every_fiscal_year_is_closed(env, buckets, app):
    from sqlalchemy.orm import Session
    from fmpoc.models import FiscalYear
    with Session(app.state.engine) as db:
        db.get(FiscalYear, buckets["base"]["fy"]["id"]).status = "CLOSED"
        db.commit()
    st = env.bu.get(url(buckets)).json()["report_status"]
    assert st["code"] == "FINAL" and st["final_because_closed"] is True


def test_cancelled_banner(env, buckets):
    assert env.bm.post(url(buckets, "/cancel"), {"reason": "Storm warning"}).status_code == 200
    t = pdf_text(get(env, buckets).content)
    assert "Cancelled: Storm warning" in t and "$1,325.00" in t


def test_access(env, buckets):
    assert get(env, buckets, client=env.admin).status_code == 403
    for c in (env.bm, env.ru, env.auditor):
        assert get(env, buckets, client=c).status_code == 200
    assert env.admin.put("/api/system/modules", {"fundraisers": False}).status_code == 200
    assert get(env, buckets).status_code == 404


def test_audit_detail_report_unchanged(env, buckets):
    """Reports -> Fundraiser detail (audit) is the former button: lines with transaction numbers and budgets."""
    r = env.auditor.get(url(buckets, "/report"))
    assert r.status_code == 200
    t = pdf_text(r.content)
    assert "Transactions" in t and "Operations" in t
    ev = env.auditor.get("/api/audit-events?action=REPORT_GENERATED").json()["items"]
    assert any((e.get("after") or {}).get("report") == "FUNDRAISER" for e in ev)


def test_long_itemized_list_uses_two_columns(env, gala):
    acct = gala["base"]["acct"]["id"]
    for i in range(14):
        env.txn(acct, "DEPOSIT", [{"budget_id": gala["base"]["inc_leaf"], "amount": f"{10 + i}.00", "description": f"Ticket {i}"}],
                date="2026-09-15")
    b = env.ru.post(url(gala, "/buckets"), {"name": "Tickets"}).json()
    bid = b["buckets"]["items"][0]["id"]
    for x in env.bu.get(url(gala)).json()["lines"]:
        if (x["description"] or "").startswith("Ticket"):
            env.ru.put(url(gala, f"/lines/{x['allocation_id']}"), {"buckets": [{"bucket_id": bid, "amount": x["amount"]}]})
    t = pdf_text(get(env, gala).content)
    assert t.count("Date Contributor Amount") >= 2      # two side-by-side tables
