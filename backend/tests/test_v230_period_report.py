"""2.3.0 build 9 (#373): the Quarterly / Annual Fiscal Year financial report and presentation.

The base Fiscal Year runs 2026-07-01 to 2027-06-30: Q1 Jul-Sep, Q2 Oct-Dec, Q3 Jan-Mar (crosses the calendar year),
Q4 Apr-Jun. "Today" is fixed through the router's `_today` (computed per request, never at import)."""
import datetime as dt
import io

import pytest
from pptx import Presentation
from pypdf import PdfReader

TODAY = dt.date(2027, 1, 20)          # Q3 in progress, Q4 not started


@pytest.fixture
def today(monkeypatch):
    import fmpoc.routers.reports as r
    box = {"d": TODAY}
    monkeypatch.setattr(r, "_today", lambda: box["d"])
    return box


def _flat(content: bytes) -> str:
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(io.BytesIO(content)).pages)


def _rep(env, fy_id, period, client=None, **kw):
    r = (client or env.auditor).get("/api/reports/period", params={"fiscal_year_id": fy_id, "period": period,
                                                                   "format": "json", **kw})
    assert r.status_code == 200, r.text
    return r.json()


def _txns(env, base):
    a = base["acct"]["id"]
    env.txn(a, "WITHDRAWAL", [{"budget_id": base["exp_leaf"], "amount": "100.00", "description": "Paper"}], date="2026-08-01")
    env.txn(a, "DEPOSIT", [{"budget_id": base["inc_leaf"], "amount": "500.00", "description": "Gift"}], date="2026-10-05")
    env.txn(a, "WITHDRAWAL", [{"budget_id": base["exp_leaf"], "amount": "50.00", "description": "Ink"}], date="2027-01-10")


def test_periods_cutoff_preliminary_and_future_blocked(env, base, today):
    fy = base["fy"]["id"]
    _txns(env, base)
    o = env.auditor.get(f"/api/reports/period/options?fiscal_year_id={fy}").json()
    sel = {p["id"]: (p["selectable"], p["preliminary"]) for p in o["periods"]}
    assert sel == {"Q1": (True, False), "Q2": (True, False), "Q3": (True, True), "Q4": (False, True),
                   "ANNUAL": (True, True)}
    r = env.auditor.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "Q4", "format": "json"})
    assert r.status_code == 422 and "hasn't started" in r.text
    # Q2 run in January: as of Dec 31, January's transaction left out (AC-RPT-002)
    q2 = _rep(env, fy, "Q2")
    assert q2["cutoff"] == "2026-12-31" and q2["preliminary"] is False
    acct = q2["accounts"][0]
    assert (acct["start"], acct["end"], acct["change"]) == (90000, 140000, 50000)
    assert [m["label"] for m in q2["cash_flow"]["months"]] == ["Oct 2026", "Nov 2026", "Dec 2026"]
    assert q2["cash_flow"]["period"] == {"label": "Q2", "in": 50000, "out": 0, "net": 50000}
    assert q2["cash_flow"]["ytd"]["out"] == 10000                      # Q1's 100, not January's 50
    # isolation: one Quarter column; Q1 only in YTD (BR-111)
    exp = q2["budgets"]["expense"]["rows"][0]
    assert q2["budgets"]["columns"] == ["Quarter"] and exp["columns"] == [0] and exp["ytd"] == 10000
    assert exp["remaining"] == 10000000 - 10000
    # the quarter in progress (crosses the calendar year): preliminary, as of today
    q3 = _rep(env, fy, "Q3")
    assert q3["cutoff"] == "2027-01-20" and q3["preliminary"] is True
    assert [m["label"] for m in q3["cash_flow"]["months"]] == ["Jan 2027"]
    assert q3["cash_flow"]["ytd"]["out"] == 15000
    # annual: Q1-Q4 columns, preliminary while the year isn't closed (AC-RPT-006)
    an = _rep(env, fy, "ANNUAL")
    assert an["preliminary"] is True and an["budgets"]["columns"] == ["Q1", "Q2", "Q3", "Q4"]
    assert an["budgets"]["expense"]["rows"][0]["columns"] == [10000, 0, 5000, 0]
    assert an["quarter_names"] == ["Q1", "Q2", "Q3"] and len(an["accounts"][0]["quarter_ends"]) == 3
    assert an["budgets"]["has_prior"] is False and an["prior_fiscal_year"] is None
    # once Q3 has ended it is final
    today["d"] = dt.date(2027, 4, 2)
    assert _rep(env, fy, "Q3")["preliminary"] is False


def test_prior_year_comparison(env, base, today):
    fy = base["fy"]["id"]
    prior = env.fy("2026", "2025-07-01", "2026-06-30")
    env.budget(prior["id"], "1000", "Operations", "EXPENSE", "900.00")
    leaf = {o["label"]: o["id"] for o in env.selectable(prior["id"], "WITHDRAWAL")}["1000 Operations"]
    a = env.account(opening="2000.00", opening_balance_date="2025-01-01")["id"]
    env.txn(a, "WITHDRAWAL", [{"budget_id": leaf, "amount": "30.00", "description": "Old paper"}], date="2025-10-10")
    env.txn(a, "WITHDRAWAL", [{"budget_id": base["exp_leaf"], "amount": "70.00", "description": "New paper"}], date="2026-11-10")
    q2 = _rep(env, fy, "Q2")
    assert q2["prior_fiscal_year"]["id"] == prior["id"] and q2["budgets"]["has_prior"]
    row = q2["budgets"]["expense"]["rows"][0]
    assert (row["prior_period"], row["prior_ytd"]) == (3000, 3000)       # matched by budget ID 1000
    assert q2["cash_flow"]["prior_period"]["out"] == 3000
    acct = next(x for x in q2["accounts"] if x["id"] == a)
    assert acct["prior_end"] == 200000 - 3000
    # income budget without a prior-year twin: no figure
    assert q2["budgets"]["income"]["rows"][0]["prior_period"] is None


def _fr(env, base, name, start, end, flt, goal=None):
    body = {"name": name, "start_date": start, "end_date": end, "filter_text": flt,
            "budget_ids": [base["inc"]["id"], base["exp"]["id"]]}
    if goal:
        body["goal_amount"] = goal
    r = env.bm.post("/api/fundraisers", body)
    assert r.status_code == 201, r.text
    return r.json()


def test_fundraiser_attribution_revised_and_cutoff(env, base, today):
    fy = base["fy"]["id"]
    assert env.admin.put("/api/system/modules", {"fundraisers": True}).status_code == 200
    a = base["acct"]["id"]
    for d, t, b, amt, desc in [("2026-09-01", "WITHDRAWAL", "exp_leaf", "200.00", "Gala supplies"),
                               ("2026-09-14", "DEPOSIT", "inc_leaf", "1200.00", "Gala tickets"),
                               ("2026-11-03", "DEPOSIT", "inc_leaf", "100.00", "Gala late gift"),
                               ("2026-11-20", "WITHDRAWAL", "exp_leaf", "50.00", "Fair supplies"),
                               ("2026-12-05", "DEPOSIT", "inc_leaf", "300.00", "Fair tickets"),
                               ("2027-01-15", "DEPOSIT", "inc_leaf", "20.00", "Fair late tickets")]:
        env.txn(a, t, [{"budget_id": base[b], "amount": amt, "description": desc}], date=d)
    gala = _fr(env, base, "Fall Gala", "2026-09-12", "2026-09-12", "Gala")
    fair = _fr(env, base, "Winter Fair", "2026-12-04", "2026-12-05", "Fair", goal="200")
    _fr(env, base, "Spring Fair", "2027-04-10", "2027-04-10", "Spring")      # ends after the cutoff: not listed
    q2 = _rep(env, fy, "Q2")
    fr = q2["fundraisers"]
    by = {r["name"]: r for r in fr["listed"]}
    assert set(by) == {"Fall Gala", "Winter Fair"}
    assert by["Fall Gala"]["revised"] and by["Fall Gala"]["change"] == 10000 and by["Fall Gala"]["quarter"] == "Q1"
    assert by["Fall Gala"]["net"] == 110000 and not by["Fall Gala"]["in_period"]
    wf = by["Winter Fair"]
    assert wf["quarter"] == "Q2" and (wf["gross"], wf["expense"], wf["net"]) == (30000, 5000, 25000)
    assert wf["after_cutoff"] is True                      # January's tickets: after Dec 31
    d = fr["details"]
    assert [x["id"] for x in d] == [fair["id"]]            # specifics only for the quarter's fundraisers
    assert d[0]["goal"]["met"] is True and d[0]["income_lines"] and d[0]["expense_lines"]
    assert q2["highlights"]["fundraiser_net"] == 25000
    # Q1 (strict cutoff): the gala's November gift isn't counted, and Q2's fair isn't there
    q1 = _rep(env, fy, "Q1")
    assert [(r["name"], r["net"], r["after_cutoff"]) for r in q1["fundraisers"]["listed"]] == [("Fall Gala", 100000, True)]
    # annual: high level only, with the fundraisers' share of income
    an = _rep(env, fy, "ANNUAL")
    assert an["fundraisers"]["details"] == [] and {r["name"] for r in an["fundraisers"]["listed"]} == {"Fall Gala", "Winter Fair"}
    assert an["fundraisers"]["total_net"] == 110000 + 27000
    assert an["fundraisers"]["share_of_income"] is not None
    assert gala["id"] in {r["id"] for r in an["fundraisers"]["rows"]}


def test_pdf_pptx_audit_and_permissions(env, base, today):
    fy = base["fy"]["id"]
    _txns(env, base)
    assert env.admin.get(f"/api/reports/period?fiscal_year_id={fy}&period=Q2").status_code == 403
    r = env.auditor.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "Q3", "commentary": "Strong giving season."})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    text = _flat(r.content)
    assert "Preliminary — as of January 20, 2027" in text and "Strong giving season." in text
    assert "Highlights" in text and "Accounts" in text and "Cash flow" in text and "Expense budgets" in text
    assert "98765432" not in text                                        # account numbers masked
    q2 = _flat(env.bm.get(f"/api/reports/period?fiscal_year_id={fy}&period=Q2").content)
    assert "Preliminary" not in q2 and "As of December 31, 2026" in q2 and "Prior" not in q2
    r = env.bm.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "ANNUAL", "format": "pptx",
                                                    "commentary": "A note for the board."})
    assert r.status_code == 200 and "presentation" in r.headers["content-type"]
    prs = Presentation(io.BytesIO(r.content))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("A note for the board." in n for n in notes)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert charts and "lineChart" in charts[0].chart._chartSpace.xml          # the net line over the monthly bars
    ev = env.auditor.get("/api/audit-events?action=REPORT_GENERATED").json()["items"]
    mine = [e for e in ev if (e.get("after") or {}).get("report") in ("FY_PERIOD_REPORT", "FY_PERIOD_PRESENTATION")]
    assert len(mine) == 3
    deck = next(e for e in mine if e["after"]["report"] == "FY_PERIOD_PRESENTATION")["after"]
    assert deck["period"] == "ANNUAL" and deck["commentary"] is True and deck["preliminary"] is True
    assert "A note for the board." not in str(deck)                     # only whether there was commentary
    # commentary is limited
    r = env.bm.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "Q2", "commentary": "x" * 2001})
    assert r.status_code == 422


# ------------------------------------------------------------------ 2.3.0 polish (Kyle, 2026-10-31)
def test_opening_balance_on_the_first_day_is_the_starting_balance(env, base, today):
    """An Opening Balance Date on the Fiscal Year's first day is the year's starting balance (BR-043), as the
    Register shows it; an account that started later has no starting balance."""
    fy = base["fy"]["id"]
    on_day = env.account(opening="250.00", opening_balance_date="2026-07-01", account_name="Savings on day one")
    later = env.account(opening="40.00", opening_balance_date="2026-07-02", account_name="Savings day two")
    env.txn(on_day["id"], "DEPOSIT", [{"budget_id": base["inc_leaf"], "amount": "10.00"}], date="2026-07-01")
    rows = {a["id"]: a for a in _rep(env, fy, "ANNUAL")["accounts"]}
    assert (rows[on_day["id"]]["start"], rows[on_day["id"]]["end"], rows[on_day["id"]]["change"]) == (25000, 26000, 1000)
    assert rows[later["id"]]["start"] is None and rows[later["id"]]["end"] == 4000
    reg = env.ru.get(f"/api/register?bank_account_id={on_day['id']}&fiscal_year_id={fy}").json()
    assert reg["starting_balance"] == "250.00"                           # the same as the Register


def test_income_above_budget_reads_as_above_budget(env, base, today):
    fy = base["fy"]["id"]
    env.txn(base["acct"]["id"], "DEPOSIT", [{"budget_id": base["inc_leaf"], "amount": "6000.00"}], date="2026-08-01")
    env.txn(base["acct"]["id"], "WITHDRAWAL", [{"budget_id": base["exp_leaf"], "amount": "100.00"}], date="2026-08-02")
    text = _flat(env.auditor.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "Q1"}).content)
    assert "+$1,000.00 above budget" in text and "-$1,000.00" not in text
    assert "$99,900.00" in text                                          # expense under budget: unchanged
    r = env.auditor.get("/api/reports/period", params={"fiscal_year_id": fy, "period": "ANNUAL", "format": "pptx"})
    prs = Presentation(io.BytesIO(r.content))
    cells = [c.text for s in prs.slides for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
    assert "+$1,000.00 above budget" in cells and "-$1,000.00" not in cells
