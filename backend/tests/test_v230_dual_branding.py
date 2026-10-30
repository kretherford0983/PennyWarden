"""2.3.0 build 10 (#386): dual branding - a light and a dark theme, the theme per type of output (new setups:
presentations dark; upgrades keep everything light), direct editing of either palette, and optional logos for dark
backgrounds. BR-003 (an Administrator setting), BR-079 (contrast for each theme), BR-100/101 (safe logo uploads)."""
from __future__ import annotations

import hashlib
import io
import sqlite3
import zipfile

from pptx import Presentation
from pypdf import PdfReader

from test_v161_fundraiser_manage import gala, put_line  # noqa: F401  (fixture)
from test_v210_fundraiser_summary import buckets  # noqa: F401  (fixture)
from test_v220_branding import URL, _png, _upload

from fmpoc.services import branding as br


def _bg(prs) -> list[str | None]:
    out = []
    for sl in prs.slides:
        f = sl.background.fill
        try:
            out.append(str(f.fore_color.rgb).lower())
        except (AttributeError, TypeError):
            out.append(None)
    return out


def test_dark_presets_meet_contrast_and_new_setups_present_dark(env, base):
    for k, p in br.DARK_PRESETS.items():
        assert all(x["ok"] for x in br.contrast_report(p["colors"], dark=True)), k
        assert len(p["series"]) == br.SERIES_COUNT and not br.similar_series(p["series"]), k
        assert set(p["colors"]) == set(br.ROLES)
    assert set(br.DARK_PRESETS) == set(br.PRESETS)
    v = env.admin.get(URL).json()
    assert v["themes"] == {"reports": "LIGHT", "presentations": "DARK"} and v["dark_logos"] is False
    assert v["dark"]["preset"] == "pennywarden" and v["dark"]["colors"] == br.DARK_PRESETS["pennywarden"]["colors"]
    assert all(x["ok"] for x in v["dark"]["contrast"])
    assert ["primary", "surface"] in v["dark_contrast_pairs"] and not any(b == "white" for _f, b in v["dark_contrast_pairs"])


def test_save_dark_palette_themes_and_guardrails(env, base):
    adm = env.admin
    # choosing a ready-made dark palette and the reports theme
    r = adm.put(URL, {"preset": "forest", "dark_preset": "ocean", "themes": {"reports": "DARK"}})
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["preset"] == "forest" and v["dark"]["preset"] == "ocean"
    assert v["themes"] == {"reports": "DARK", "presentations": "DARK"}
    # a custom dark palette: every role, eight series, WCAG AA on the dark background
    dark = dict(br.DARK_PRESETS["plum"]["colors"])
    series = list(br.DARK_PRESETS["plum"]["series"])
    bad = {**dark, "on_surface": "#303030"}
    r = adm.put(URL, {"preset": "forest", "dark_preset": "custom", "dark_colors": bad, "dark_series": series})
    assert r.status_code == 422 and "Dark theme" in r.text and "dark_colors.on_surface" in r.text
    r = adm.put(URL, {"preset": "forest", "dark_preset": "custom", "dark_colors": {**dark, "primary": "#e8d0ff"},
                      "dark_series": series, "themes": {"presentations": "LIGHT"}, "dark_logos": True})
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["dark"]["preset"] == "custom" and v["dark"]["colors"]["primary"] == "#e8d0ff"
    assert v["themes"]["presentations"] == "LIGHT" and v["dark_logos"] is True
    assert adm.put(URL, {"preset": "forest", "dark_preset": "custom", "dark_colors": dark,
                         "dark_series": series[:7]}).status_code == 422
    assert adm.put(URL, {"preset": "forest", "themes": {"reports": "SEPIA"}}).status_code == 422
    assert adm.put(URL, {"preset": "forest", "themes": {"letters": "DARK"}}).status_code == 422
    # similar dark chart colors need the confirmation
    close = [series[0]] * 2 + series[2:]
    r = adm.put(URL, {"preset": "forest", "dark_preset": "custom", "dark_colors": dark, "dark_series": close})
    assert r.status_code == 409 and "dark theme 1 and 2" in r.text
    assert adm.put(URL, {"preset": "forest", "dark_preset": "custom", "dark_colors": dark, "dark_series": close,
                         "confirmations": ["SIMILAR_SERIES_COLORS"]}).status_code == 200
    # leaving the dark fields out keeps them; changes are audited with before/after
    v = adm.put(URL, {"preset": "ember"}).json()
    assert v["dark"]["preset"] == "custom" and v["themes"]["reports"] == "DARK"
    ev = env.auditor.get("/api/audit-events?action=BRANDING_UPDATED").json()["items"]
    assert any(e["after"].get("dark_preset") == "ocean" for e in ev)
    # reset: both palettes back to PennyWarden; the themes per output and the logo choice stay
    v = adm.post(f"{URL}/reset", {}).json()
    assert v["preset"] == "pennywarden" and v["dark"]["preset"] == "pennywarden"
    assert v["themes"]["reports"] == "DARK" and v["dark_logos"] is True
    # Administrators only (BR-003; financial users see branding read-only elsewhere)
    for c in (env.bm, env.bu, env.auditor):
        assert c.put(URL, {"preset": "forest", "themes": {"reports": "LIGHT"}}).status_code == 403
        assert _upload(c, "dark_full", _png()).status_code == 403


def test_dark_logos_are_safe_uploads(env, base):
    adm = env.admin
    r = _upload(adm, "dark_full", _png(meta=True))
    assert r.status_code == 200, r.text
    assert r.json()["logos"]["dark_full"]["width"] == 400
    png = adm.get(f"{URL}/logo/dark_full").content
    assert b"tEXt" not in png                               # re-encoded without metadata (BR-101)
    assert _upload(adm, "dark_mark", b"<svg/>", name="x.svg", ctype="image/svg+xml").status_code in (415, 422)
    assert _upload(adm, "dark_banner", _png()).status_code == 422
    assert adm.post(f"{URL}/logo/dark_full/remove", {}).json()["logos"]["dark_full"] is None


def test_presentations_dark_reports_light_word_light(env, base, buckets):  # noqa: F811
    g = buckets
    adm = env.admin
    assert adm.put(URL, {"preset": "plum", "dark_preset": "plum"}).status_code == 200
    light_logo, dark_logo = _png(color=(10, 120, 30)), _png(color=(240, 240, 240))
    assert _upload(adm, "full", light_logo).status_code == 200
    assert _upload(adm, "dark_full", dark_logo).status_code == 200
    surface = br.DARK_PRESETS["plum"]["colors"]["surface"][1:]
    # the deck: dark slides, the dark palette, the light logo while separate dark logos are off
    prs = Presentation(io.BytesIO(env.bu.get(f"/api/fundraisers/{g['id']}/presentation").content))
    assert set(_bg(prs)) == {surface}
    xml = b"".join(p.blob for p in prs.part.package.iter_parts() if p.partname.endswith(".xml"))
    assert br.DARK_PRESETS["plum"]["colors"]["primary"][1:].upper().encode() in xml
    pics = [sh for sh in list(prs.slides)[0].shapes if sh.shape_type == 13]
    light_sha = hashlib.sha256(adm.get(f"{URL}/logo/full").content).hexdigest()
    dark_sha = hashlib.sha256(adm.get(f"{URL}/logo/dark_full").content).hexdigest()
    assert hashlib.sha256(pics[0].image.blob).hexdigest() == light_sha
    # ... and the dark logo once they are on
    assert adm.put(URL, {"preset": "plum", "dark_logos": True}).status_code == 200
    prs = Presentation(io.BytesIO(env.bu.get(f"/api/fundraisers/{g['id']}/presentation").content))
    pics = [sh for sh in list(prs.slides)[0].shapes if sh.shape_type == 13]
    assert hashlib.sha256(pics[0].image.blob).hexdigest() == dark_sha
    # printed reports stay light (no page fill), Word always light
    pdf = env.bu.get(f"/api/fundraisers/{g['id']}/summary-report").content
    assert b"0 0 612 792 re" not in b"".join(p.get_contents().get_data() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert adm.put(URL, {"preset": "plum", "themes": {"reports": "DARK"}}).status_code == 200
    pdf = env.bu.get(f"/api/fundraisers/{g['id']}/summary-report").content
    first = PdfReader(io.BytesIO(pdf)).pages[0].get_contents().get_data()
    assert b"0 0 612 792 re" in first                     # the dark surface under the page
    d = env.bu.get(f"/api/fundraisers/{g['id']}/summary-report?format=docx").content
    with zipfile.ZipFile(io.BytesIO(d)) as z:
        doc = z.read("word/document.xml") + b"".join(z.read(n) for n in z.namelist() if n.startswith("word/styles"))
    assert br.PRESETS["plum"]["colors"]["primary"][1:].upper().encode() in doc.upper()
    # a light presentation theme gives white slides again
    assert adm.put(URL, {"preset": "plum", "themes": {"presentations": "LIGHT"}}).status_code == 200
    prs = Presentation(io.BytesIO(env.bu.get(f"/api/fundraisers/{g['id']}/presentation").content))
    assert surface not in _bg(prs)


def test_upgrade_keeps_existing_installations_light(tmp_path):
    """0065: every existing organization keeps light for every output type (a branding row is added where none
    existed); only new installations start with dark presentations."""
    from alembic import command

    from fmpoc.db import alembic_config
    url = f"sqlite:///{tmp_path / 'm.sqlite3'}"
    command.upgrade(alembic_config(url), "0064_shared_signers")
    c = sqlite3.connect(tmp_path / "m.sqlite3")
    c.execute("pragma foreign_keys=off")
    for ws in (1, 2):
        cols = [r for r in c.execute("pragma table_info(workspace)").fetchall()]
        row = {r[1]: None for r in cols}
        row.update(id=ws, name=f"Org {ws}", created_at="2026-01-01 00:00:00")
        for r in cols:
            if r[3] and row.get(r[1]) is None:
                row[r[1]] = 0 if "INT" in (r[2] or "").upper() or "BOOL" in (r[2] or "").upper() else ""
        c.execute(f"insert into workspace ({','.join(row)}) values ({','.join('?' * len(row))})", list(row.values()))
    c.execute("insert into branding_settings (workspace_id, preset) values (1, 'forest')")
    c.commit()
    c.close()
    command.upgrade(alembic_config(url), "head")
    c = sqlite3.connect(tmp_path / "m.sqlite3")
    rows = {r[0]: r[1:] for r in c.execute(
        "select workspace_id, preset, theme_reports, theme_presentations, dark_logos, dark_preset from branding_settings")}
    assert rows == {1: ("forest", "LIGHT", "LIGHT", 0, None), 2: ("pennywarden", "LIGHT", "LIGHT", 0, None)}
    c.close()
